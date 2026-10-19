"""Uploading Marketing's PPC sheet ("SaaS PPC Campaigns") instead of retyping it.

Preview → confirm → (optionally) undo, like the cost sheet upload (services/cost_sheet_import_service.py).
The preview matches each row to a Google Ads or Bing Ads campaign (an initiative), or marks it as a new campaign,
and says whether it would add a result, update one, or change nothing. Confirming creates the new campaigns and
upserts the results through ads_service.apply_results (source SPREADSHEET), then syncs spend and outcomes.
Undo removes the results the upload created and gives the ones it updated their previous values back.
"""
import io
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from uuid import UUID

from postgresql_db.database import get_pool
from services import ads_service
from services.cost_sheet_import_service import (
    ImportConflictError,
    ImportExpiredError,
    ImportForbiddenError,
    ImportNotFoundError,
)
from services.import_parsing import (
    ROW_NUMBER,
    ImportFileError,
    InvalidValueError,
    match_columns,
    normalise_header,
    parse_amount,
    parse_date,
    read_table,
    uses_decimal_comma,
)
from services.rbac_service import log_audit

KIND = "ADS_RESULTS"
AUDIT_ENTITY = "ImportBatch"
INITIATIVE_AUDIT_ENTITY = "Initiative"
MAX_FILE_BYTES = 5 * 1024 * 1024
PREVIEW_LIFETIME = timedelta(hours=24)
SHEET_NAMES = ["SaaS PPC Campaigns"]

# Sheet header → our field (compared case-insensitively; spaces, underscores and dashes count the same).
# CTR, CPC, Ads Conv Rate, CPL and Last Update aren't read: the portal calculates them. Primary KPI
# stays in the template only because it's Marketing's own layout: it's ignored, no KPI is set on campaigns.
COLUMN_ALIASES: dict[str, list[str]] = {
    "brand": ["Brand"],
    "market": ["Market/Region", "Market", "Region"],
    "period_start": ["Period Start", "Start", "Start date"],
    "period_end": ["Period End", "End", "End date"],
    "channel": ["Channel"],
    "sub_channel": ["Sub-Channel", "Sub channel", "Platform"],
    "campaign": ["Campaign Name", "Campaign"],
    "ad_group": ["Ad Group", "Ad group name"],
    "budget_amount": ["Budget Allocated ($)", "Budget Allocated", "Budget"],
    "spend": ["Actual Spend", "Spend", "Cost"],
    "impressions": ["Impressions", "Impr."],
    "clicks": ["Clicks"],
    "conversions": ["Ads Conv.", "Ads Conv", "Conversions", "Conv."],
    "hubspot_contacts": ["HS contact", "HS contacts", "HubSpot contacts"],
    "mqls": ["MQLs", "MQL"],
    "sqls": ["SQLs", "SQL"],
    "status": ["Status"],
    "report_url": ["Report Link", "Report", "Report URL"],
    "tracking_detail": ["Tracking Detail", "Tracking"],
}
# Marketing's own header, in order, so they can keep their layout
TEMPLATE_COLUMNS = [
    "Brand", "Market/Region", "Period Start", "Period End", "Channel", "Sub-Channel", "Campaign Name", "Primary KPI",
    "Budget Allocated ($)", "Actual Spend", "Impressions", "Clicks", "CTR (%)", "CPC ($)", "Ads Conv.", "HS contact",
    "Ads Conv Rate", "CPL ($)", "MQLs", "SQLs", "Last Update", "Status", "Report Link", "Tracking Detail",
]

_PLATFORM_WORDS = {
    "google ads": "GOOGLE_ADS", "google": "GOOGLE_ADS", "adwords": "GOOGLE_ADS",
    "bing ads": "BING_ADS", "bing": "BING_ADS", "microsoft ads": "BING_ADS", "microsoft advertising": "BING_ADS",
}
_STATUS_WORDS = {
    "ongoing": "ACTIVE", "active": "ACTIVE", "live": "ACTIVE", "running": "ACTIVE", "enabled": "ACTIVE",
    "paused": "PAUSED", "on hold": "PAUSED",
    "ended": "ENDED", "complete": "ENDED", "completed": "ENDED", "stopped": "ENDED", "finished": "ENDED",
}
# Cells that mean "nothing here": spreadsheet errors from the calculated columns, and placeholders
_BLANK_CELLS = {"", ".", "-", "–", "n/a", "#n/a", "#ref!", "#value!", "#div/0!", "#name?", "#num!", "#null!"}
_COUNT_FIELDS = {
    "impressions": "Impressions", "clicks": "Clicks", "hubspot_contacts": "HS contact", "mqls": "MQLs", "sqls": "SQLs",
}
_AMOUNT_FIELDS = {"budget_amount": "Budget Allocated", "spend": "Actual Spend", "conversions": "Ads Conv."}
_FIELD_LABELS = {
    "period_end": "period end", "budget_amount": "budget", "spend": "spend", "impressions": "impressions",
    "clicks": "clicks", "conversions": "conversions", "hubspot_contacts": "HS contacts", "mqls": "MQLs",
    "sqls": "SQLs", "status": "status", "report_url": "report link", "tracking_detail": "tracking detail",
}
_MAX_AMOUNT = Decimal("999999999999")
_MAX_COUNT = {"impressions": 9_000_000_000_000_000, "clicks": 9_000_000_000_000_000}
_MAX_INT = 2_147_483_647


def _join_words(words: list[str]) -> str:
    return words[0] if len(words) == 1 else f"{', '.join(words[:-1])} and {words[-1]}"


# --- Checking one row ---

class _Context:
    """Everything a row check needs, loaded once per upload."""

    def __init__(self, columns: dict[str, str], decimal_comma: bool, campaigns, brands, markets):
        self.columns = columns
        self.decimal_comma = decimal_comma
        # Per platform: campaign names (lower case) → initiatives with that name. A Bing campaign also answers
        # to its name without " (Bing)".
        self.by_name: dict[str, dict[str, list[dict]]] = {p: {} for p in ads_service.PLATFORMS}
        for c in campaigns:
            campaign = dict(c)
            names = {campaign["name"].strip().lower()}
            if campaign["platform"] == "BING_ADS" and campaign["name"].lower().endswith(" (bing)"):
                names.add(campaign["name"][: -len(" (bing)")].strip().lower())
            for name in names:
                self.by_name[campaign["platform"]].setdefault(name, []).append(campaign)
        # Longest first, so "LP1 - Brand" wins over "LP1" when both are prefixes
        self.prefixes = {p: sorted(names, key=len, reverse=True) for p, names in self.by_name.items()}
        self.brands = _lookup(brands)
        self.markets = _lookup(markets)


def _lookup(rows) -> dict[str, dict]:
    found: dict[str, dict] = {}
    for r in rows:
        found.setdefault(normalise_header(r["name"]), dict(r))
        found.setdefault(normalise_header(r["code"]), dict(r))
    return found


def _cell(row: dict, ctx: _Context, field: str) -> str:
    header = ctx.columns.get(field)
    return str(row.get(header, "")).strip() if header else ""


def _is_blank(text: str) -> bool:
    return text.strip().lower() in _BLANK_CELLS


def _match_campaign(ctx: _Context, platform: str, name: str, ad_group_cell: str):
    """(initiative, ad group, problem) for a campaign name; initiative None when it's a new campaign."""
    key = name.lower()
    exact = ctx.by_name[platform].get(key)
    if exact:
        if len(exact) > 1:
            return None, None, f"More than one {ads_service.PLATFORM_LABELS[platform]} campaign is called '{name}'. Rename one in the portal."
        return exact[0], ad_group_cell or None, None
    if not ad_group_cell:
        for prefix in ctx.prefixes[platform]:
            if key.startswith(prefix + " - ") and len(ctx.by_name[platform][prefix]) == 1:
                return ctx.by_name[platform][prefix][0], name[len(prefix) + 3:].strip() or None, None
    return None, ad_group_cell or None, None


def _check_row(row: dict, ctx: _Context) -> dict:
    """Read and check one sheet row. Returns the preview row plus what would be imported (under '_staged')."""
    errors: list[str] = []
    warnings: list[str] = []

    raw_platform = _cell(row, ctx, "sub_channel")
    platform = _PLATFORM_WORDS.get(normalise_header(raw_platform)) if raw_platform else None
    if not raw_platform:
        errors.append("Add the sub-channel (Google Ads or Bing Ads).")
    elif platform is None:
        errors.append("Only Google Ads and Bing Ads rows can be imported here.")

    campaign_name = " ".join(_cell(row, ctx, "campaign").split())
    if not campaign_name:
        errors.append("Add the campaign name.")
    elif len(campaign_name) > 255:
        errors.append("Campaign name is longer than 255 characters.")
    ad_group_cell = " ".join(_cell(row, ctx, "ad_group").split())
    if _is_blank(ad_group_cell):
        ad_group_cell = ""

    raw_brand = _cell(row, ctx, "brand")
    brand = ctx.brands.get(normalise_header(raw_brand)) if raw_brand else None
    if not raw_brand:
        errors.append("Add the brand.")
    elif brand is None:
        errors.append(f"Unknown brand '{raw_brand}'.")
    raw_market = _cell(row, ctx, "market")
    market = ctx.markets.get(normalise_header(raw_market)) if raw_market else None
    if not raw_market:
        errors.append("Add the market or region.")
    elif market is None:
        errors.append(f"Unknown market '{raw_market}'.")

    def day(field: str, label: str) -> date | None:
        try:
            return parse_date(_cell(row, ctx, field))
        except InvalidValueError as exc:
            errors.append(f"{label}: {exc}")
            return None

    period_start = day("period_start", "Period Start")
    if period_start is None and not _cell(row, ctx, "period_start"):
        errors.append("Add the period start.")
    period_end = day("period_end", "Period End")
    if period_start and period_end and period_end < period_start:
        errors.append("Period End is before Period Start.")

    values: dict = {}
    for field, label in {**_AMOUNT_FIELDS, **_COUNT_FIELDS}.items():
        text = _cell(row, ctx, field)
        if _is_blank(text):
            values[field] = None
            continue
        try:
            number = parse_amount(text, decimal_comma=ctx.decimal_comma)
        except InvalidValueError as exc:
            errors.append(f"{label}: {exc}")
            values[field] = None
            continue
        if number is None:
            values[field] = None
        elif number < 0:
            errors.append(f"{label} can't be negative.")
            values[field] = None
        elif field in _COUNT_FIELDS:
            if number != number.to_integral_value():
                errors.append(f"{label} must be a whole number.")
                values[field] = None
            elif number > _MAX_COUNT.get(field, _MAX_INT):
                errors.append(f"{label} is too large.")
                values[field] = None
            else:
                values[field] = int(number)
        elif number > _MAX_AMOUNT:
            errors.append(f"{label} is too large.")
            values[field] = None
        else:
            values[field] = number.quantize(Decimal("0.01"))

    raw_status = _cell(row, ctx, "status")
    status = "ACTIVE"
    if not _is_blank(raw_status):
        status = _STATUS_WORDS.get(normalise_header(raw_status))
        if status is None:
            errors.append(f"Unknown status '{raw_status}'. Use ongoing, paused or ended.")

    report_url = _cell(row, ctx, "report_url")
    if _is_blank(report_url):
        report_url = None
    elif not report_url.lower().startswith("https://") or len(report_url) > 2000 or any(c.isspace() for c in report_url):
        warnings.append("The report link isn't an https link, so it was left out.")
        report_url = None
    tracking = _cell(row, ctx, "tracking_detail")
    tracking = None if _is_blank(tracking) else tracking[:2000]

    initiative, ad_group = None, ad_group_cell or None
    if platform and campaign_name:
        initiative, ad_group, problem = _match_campaign(ctx, platform, campaign_name, ad_group_cell)
        if problem:
            errors.append(problem)
    if ad_group and len(ad_group) > 255:
        errors.append("Ad group is longer than 255 characters.")

    preview = {
        "row_number": row[ROW_NUMBER],
        "status": "INVALID" if errors else "NEW",
        "messages": errors + warnings,
        "campaign_name": initiative["name"] if initiative else (campaign_name or None),
        "ad_group": ad_group,
        "platform": platform,
        "brand": brand["name"] if brand else raw_brand or None,
        "market": market["name"] if market else raw_market or None,
        "period_start": period_start,
        "period_end": period_end,
        **{f: (float(v) if isinstance(v, Decimal) else v) for f, v in values.items()},
        "result_status": status,
        "currency": initiative["currency"] if initiative else ("USD" if platform and campaign_name else None),
        "report_url": report_url,
        "initiative_id": initiative["id"] if initiative else None,
        "_staged": None,
    }
    if not errors:
        preview["_staged"] = {
            "row_number": row[ROW_NUMBER],
            "initiative_id": initiative["id"] if initiative else None,
            "campaign_name": campaign_name,
            "platform": platform,
            "brand_id": brand["id"],
            "market_id": market["id"],
            "ad_group": ad_group,
            "period_start": period_start,
            "period_end": period_end,
            **values,
            "status": status,
            "report_url": report_url,
            "tracking_detail": tracking,
        }
    return preview


def _link_new_ad_groups(previews: list[dict]) -> None:
    """A row like "X - LP9 - Ad Group 1a" whose campaign "X - LP9" is new in this file is an ad group of it."""
    new_names: dict[str, set[str]] = {}
    for p in previews:
        staged = p["_staged"]
        if staged and staged["initiative_id"] is None and not staged["ad_group"]:
            new_names.setdefault(staged["platform"], set()).add(staged["campaign_name"].lower())
    for p in previews:
        staged = p["_staged"]
        if not staged or staged["initiative_id"] is not None or staged["ad_group"]:
            continue
        name = staged["campaign_name"]
        for prefix in sorted(new_names.get(staged["platform"], ()), key=len, reverse=True):
            if name.lower().startswith(prefix + " - "):
                staged["ad_group"] = p["ad_group"] = name[len(prefix) + 3:].strip()
                staged["campaign_name"] = p["campaign_name"] = name[: len(prefix)]
                break


def _new_campaign_key(staged: dict) -> tuple[str, str]:
    return staged["platform"], staged["campaign_name"].lower()


def _json_ready(staged: dict) -> dict:
    return {k: (v.isoformat() if isinstance(v, date) else str(v) if isinstance(v, (Decimal, UUID)) else v) for k, v in staged.items()}


# --- Preview ---

async def preview_ads_sheet(filename: str, data: bytes, actor_id: UUID) -> dict:
    """Check every row of an uploaded PPC sheet and keep the rows that pass for confirming later."""
    known = [name for names in COLUMN_ALIASES.values() for name in names]
    rows = read_table(filename, data, known, sheet_names=SHEET_NAMES)
    if not rows:
        raise ImportFileError("The sheet has a header but no campaign rows.")
    headers = [h for h in rows[0] if h != ROW_NUMBER]
    match = match_columns(headers, COLUMN_ALIASES)
    labels = (("sub_channel", "Sub-Channel"), ("campaign", "Campaign Name"), ("period_start", "Period Start"))
    missing = [label for field, label in labels if field not in match.fields]
    if missing:
        plural = "s" if len(missing) > 1 else ""
        raise ImportFileError(f"The sheet needs {_join_words(missing)} column{plural}. Start from the template.")
    number_cells = [
        row.get(match.fields[f], "") for row in rows for f in (*_AMOUNT_FIELDS, *_COUNT_FIELDS) if f in match.fields
    ]

    pool = get_pool()
    async with pool.acquire() as conn:
        # A fixed number of queries, however long the sheet is
        campaigns = await conn.fetch(f"SELECT i.id, i.name, i.currency, sc.code AS platform {ads_service.CAMPAIGNS_FROM}")
        brands = await conn.fetch("SELECT id, code, name FROM brands WHERE is_active")
        markets = await conn.fetch("SELECT id, code, name FROM markets WHERE is_active")
        ctx = _Context(match.fields, uses_decimal_comma(number_cells), campaigns, brands, markets)
        previews = [_check_row(row, ctx) for row in rows]
        _link_new_ad_groups(previews)

        matched = [p["_staged"] for p in previews if p["_staged"] and p["_staged"]["initiative_id"]]
        existing_rows = await conn.fetch(
            f"""
            SELECT r.initiative_id, r.ad_group, r.period_start, {', '.join('r.' + c for c in ads_service.FIGURE_COLUMNS)}
            FROM ad_results r WHERE r.initiative_id = ANY($1::uuid[])
            """,
            sorted({s["initiative_id"] for s in matched}, key=str),
        )
        existing = {(r["initiative_id"], r["ad_group"] or "", r["period_start"]): dict(r) for r in existing_rows}

        seen: dict[tuple, int] = {}
        staged_rows = []
        for p in previews:
            staged = p.pop("_staged")
            if staged is None:
                continue
            campaign_key = staged["initiative_id"] or _new_campaign_key(staged)
            key = (campaign_key, (staged["ad_group"] or "").lower(), staged["period_start"])
            if key in seen:
                p["status"] = "INVALID"
                p["messages"] = [f"Appears twice in this file (row {seen[key]})."]
                continue
            seen[key] = p["row_number"]
            if staged["initiative_id"] is None:
                p["status"] = "NEW_CAMPAIGN"
                label = ads_service.PLATFORM_LABELS[staged["platform"]]
                p["messages"] = [f"A new {label} campaign called '{staged['campaign_name']}' will be added.", *p["messages"]]
            else:
                current = existing.get((staged["initiative_id"], staged["ad_group"] or "", staged["period_start"]))
                if current is not None:
                    changed = ads_service.changed_figures(current, staged)
                    if changed:
                        p["status"] = "UPDATE"
                        words = [_FIELD_LABELS[c] for c in ads_service.FIGURE_COLUMNS if c in changed]
                        p["messages"] = [f"Updates {_join_words(words)}.", *p["messages"]]
                    else:
                        p["status"] = "UNCHANGED"
            # Unchanged rows are kept too: they may have changed by the time the upload is confirmed
            staged_rows.append(_json_ready(staged))

        counts = {
            "new": sum(p["status"] == "NEW" for p in previews),
            "update": sum(p["status"] == "UPDATE" for p in previews),
            "unchanged": sum(p["status"] == "UNCHANGED" for p in previews),
            "invalid": sum(p["status"] == "INVALID" for p in previews),
            "new_campaign": sum(p["status"] == "NEW_CAMPAIGN" for p in previews),
        }
        batch_id = await conn.fetchval(
            """
            INSERT INTO import_batches (kind, filename, status, total_rows, ready_rows, error_rows, duplicate_rows,
                                        staged_rows, created_by)
            VALUES ($1, $2, 'PREVIEW', $3, $4, $5, $6, $7::jsonb, $8)
            RETURNING id
            """,
            KIND, (filename or "upload")[:255], len(previews), len(staged_rows) - counts["unchanged"], counts["invalid"],
            0, json.dumps(staged_rows), actor_id,
        )
    return {"batch_id": batch_id, "filename": filename, "rows": previews, "counts": counts}


# --- Confirm and undo ---

def _is_expired(created_at: datetime) -> bool:
    return created_at < datetime.now(timezone.utc) - PREVIEW_LIFETIME


def _from_json(staged: dict) -> dict:
    row = dict(staged)
    for field in ("period_start", "period_end"):
        row[field] = date.fromisoformat(row[field]) if row.get(field) else None
    for field in _AMOUNT_FIELDS:
        row[field] = Decimal(row[field]) if row.get(field) is not None else None
    row["initiative_id"] = UUID(row["initiative_id"]) if row.get("initiative_id") else None
    return row


async def _create_campaigns(conn, staged: list[dict], actor_id: UUID) -> list[UUID]:
    """Create the new campaigns the rows need (or reuse one added since the preview); fill in initiative_id.

    Returns the ids of the campaigns created.
    """
    wanted: dict[tuple[str, str], dict] = {}
    for row in sorted(staged, key=lambda r: r["row_number"]):
        if row["initiative_id"] is None:
            key = _new_campaign_key(row)
            campaign = wanted.setdefault(key, {**row, "start_date": row["period_start"]})
            campaign["start_date"] = min(campaign["start_date"], row["period_start"])
    if not wanted:
        return []

    found = await conn.fetch(
        f"SELECT i.id, sc.code AS platform, lower(i.name) AS name_key {ads_service.CAMPAIGNS_FROM} "
        "WHERE lower(i.name) = ANY($1::text[]) ORDER BY i.created_at",
        sorted({name for _, name in wanted}),
    )
    ids: dict[tuple[str, str], UUID] = {}
    for r in found:
        ids.setdefault((r["platform"], r["name_key"]), r["id"])
    to_create = [(key, c) for key, c in wanted.items() if key not in ids]
    created = []
    if to_create:
        refs = await conn.fetchrow(
            "SELECT (SELECT id FROM channels WHERE code = 'PAID_SEARCH') AS channel_id, "
            "(SELECT id FROM sub_channels WHERE code = 'GOOGLE_ADS') AS google, "
            "(SELECT id FROM sub_channels WHERE code = 'BING_ADS') AS bing"
        )
        sub_channels = {"GOOGLE_ADS": refs["google"], "BING_ADS": refs["bing"]}
        created = await conn.fetch(
            """
            INSERT INTO initiatives (name, brand_id, channel_id, sub_channel_id, owner_user_id, start_date, currency,
                                     status, market_id, created_by, updated_by)
            SELECT u.name, u.brand_id, $1, u.sub_channel_id, $2, u.start_date, 'USD', 'ACTIVE', u.market_id, $2, $2
            FROM unnest($3::text[], $4::int[], $5::int[], $6::date[], $7::int[])
                 WITH ORDINALITY AS u(name, brand_id, sub_channel_id, start_date, market_id, n)
            ORDER BY u.n
            RETURNING id, name, brand_id, sub_channel_id, start_date, market_id
            """,
            refs["channel_id"], actor_id,
            [c["campaign_name"] for _, c in to_create], [c["brand_id"] for _, c in to_create],
            [sub_channels[c["platform"]] for _, c in to_create], [c["start_date"] for _, c in to_create],
            [c["market_id"] for _, c in to_create],
        )
        audits = []
        for (key, _), row in zip(to_create, created):
            ids[key] = row["id"]
            audits.append(("CREATE", row["id"], None, {
                **{k: row[k] for k in ("name", "brand_id", "sub_channel_id", "start_date", "market_id")},
                "channel_id": refs["channel_id"], "currency": "USD", "status": "ACTIVE",
                "owner_user_id": actor_id, "import": "PPC sheet",
            }))
        await ads_service.audit_many(conn, actor_id, INITIATIVE_AUDIT_ENTITY, audits)
    for row in staged:
        if row["initiative_id"] is None:
            row["initiative_id"] = ids[_new_campaign_key(row)]
    return [r["id"] for r in created]


async def confirm_import(batch_id: UUID, actor_id: UUID) -> dict:
    """Create the new campaigns and import the previewed rows, then mark the upload confirmed."""
    expired = False
    pool = get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            batch = await conn.fetchrow(
                "SELECT kind, filename, status, staged_rows, created_by, created_at FROM import_batches "
                "WHERE id = $1 AND kind = $2 FOR UPDATE",
                batch_id, KIND,
            )
            if batch is None:
                raise ImportNotFoundError
            if batch["created_by"] != actor_id:
                raise ImportForbiddenError
            if batch["status"] != "PREVIEW":
                raise ImportConflictError(f"This upload is already {batch['status'].lower()}.")
            if _is_expired(batch["created_at"]):
                await conn.execute("UPDATE import_batches SET status = 'EXPIRED' WHERE id = $1", batch_id)
                expired = True
            else:
                raw = json.loads(batch["staged_rows"]) if isinstance(batch["staged_rows"], str) else batch["staged_rows"]
                staged = [_from_json(r) for r in raw]

                # Campaigns removed (or moved out of Google Ads / Bing Ads) since the preview are skipped
                matched_ids = sorted({r["initiative_id"] for r in staged if r["initiative_id"]}, key=str)
                still_ads = {
                    r["id"] for r in await conn.fetch(
                        f"SELECT i.id {ads_service.CAMPAIGNS_FROM} WHERE i.id = ANY($1::uuid[])", matched_ids
                    )
                }
                usable = [r for r in staged if r["initiative_id"] is None or r["initiative_id"] in still_ads]
                skipped = len(staged) - len(usable)

                created_campaign_ids = await _create_campaigns(conn, usable, actor_id)
                try:
                    applied = await ads_service.apply_results(conn, usable, "SPREADSHEET", actor_id, batch_id)
                except ads_service.AdsValidationError as exc:
                    raise ImportConflictError("Some rows no longer pass the checks. Upload the file again.") from exc
                conflicts = await ads_service.sync_initiatives(conn, [r["initiative_id"] for r in usable], actor_id)
                undo_data = {
                    "created": [str(i) for i in applied["created"]],
                    "updated": applied["updated"],
                    "campaigns": [str(i) for i in created_campaign_ids],
                }
                created, updated = len(applied["created"]), len(applied["updated"])
                await conn.execute(
                    """
                    UPDATE import_batches
                    SET status = 'CONFIRMED', imported_rows = $2, confirmed_at = now(), confirmed_by = $3,
                        undo_data = $4::jsonb
                    WHERE id = $1
                    """,
                    batch_id, created + updated, actor_id, json.dumps(undo_data, default=str),
                )
                summary = {
                    "created": created, "updated": updated, "unchanged": applied["unchanged"],
                    "campaigns_created": len(created_campaign_ids), "skipped": skipped, "spend_conflicts": conflicts,
                }
                await log_audit(
                    conn, actor_id, "IMPORT", AUDIT_ENTITY, batch_id,
                    new_value={"kind": KIND, "filename": batch["filename"], **summary},
                )
    if expired:
        raise ImportExpiredError
    return {"batch_id": batch_id, "status": "CONFIRMED", **summary}


_RESTORE_TYPES = {**{c: ads_service._SQL_TYPES[c] for c in ads_service.FIGURE_COLUMNS}, "source": "text", "import_batch_id": "uuid"}


def _restore_value(column: str, value):
    if value is None:
        return None
    kind = _RESTORE_TYPES[column]
    if kind == "date":
        return date.fromisoformat(value)
    if kind == "numeric":
        return Decimal(str(value))
    if kind in ("bigint", "int"):
        return int(value)
    if kind == "uuid":
        return UUID(str(value))
    return value


async def undo_import(batch_id: UUID, actor_id: UUID) -> dict:
    """Remove the results an upload created and restore the previous values of the ones it updated.

    Rows a later upload changed again, and rows whose cost Finance already reconciled, are left as they are.
    Campaigns the upload created are removed too when nothing else uses them.
    """
    pool = get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            batch = await conn.fetchrow(
                "SELECT filename, status, undo_data FROM import_batches WHERE id = $1 AND kind = $2 FOR UPDATE",
                batch_id, KIND,
            )
            if batch is None:
                raise ImportNotFoundError
            if batch["status"] != "CONFIRMED":
                raise ImportConflictError("Only a confirmed upload can be undone.")
            data = batch["undo_data"]
            data = json.loads(data) if isinstance(data, str) else (data or {})
            created_ids = [UUID(i) for i in data.get("created", [])]
            previous = {UUID(u["id"]): u["previous"] for u in data.get("updated", [])}

            # Rows still as this upload left them, and whose cost isn't reconciled
            candidates = await conn.fetch(
                f"""
                SELECT r.id, r.initiative_id, i.name AS campaign_name, r.ad_group, r.period_start,
                       {', '.join('r.' + c for c in ads_service.FIGURE_COLUMNS)}
                FROM ad_results r
                JOIN initiatives i ON i.id = r.initiative_id
                WHERE r.id = ANY($1::uuid[]) AND r.import_batch_id = $2
                  AND NOT EXISTS (
                      SELECT 1 FROM spend_lines s
                      WHERE s.external_ref = '{ads_service.EXTERNAL_REF_PREFIX}' || r.id::text
                        AND s.source IN ({ads_service.AD_SOURCES_SQL}) AND s.reconciled_month IS NOT NULL
                  )
                FOR UPDATE OF r
                """,
                [*created_ids, *previous], batch_id,
            )
            removing = [c for c in candidates if c["id"] not in previous]
            restoring = [c for c in candidates if c["id"] in previous]
            audits = []
            if removing:
                await conn.execute("DELETE FROM ad_results WHERE id = ANY($1::uuid[])", [c["id"] for c in removing])
                audits += [("DELETE", c["id"], ads_service._audit_row(dict(c)), None) for c in removing]
            if restoring:
                columns = list(_RESTORE_TYPES)
                params = ", ".join(f"${n}::{_RESTORE_TYPES[c]}[]" for n, c in enumerate(columns, start=3))
                await conn.execute(
                    f"""
                    UPDATE ad_results r
                    SET {', '.join(f'{c} = u.{c}' for c in columns)}, updated_at = now(), updated_by = $1
                    FROM unnest($2::uuid[], {params}) AS u(id, {', '.join(columns)})
                    WHERE r.id = u.id
                    """,
                    actor_id, [c["id"] for c in restoring],
                    *[[_restore_value(col, previous[c["id"]].get(col)) for c in restoring] for col in columns],
                )
                for c in restoring:
                    old = {k: c[k] for k in ("initiative_id", "campaign_name", *ads_service.FIGURE_COLUMNS)}
                    audits.append(("UPDATE", c["id"], ads_service._audit_row(old),
                                   {"initiative_id": str(c["initiative_id"]), **previous[c["id"]]}))
            await ads_service.audit_many(conn, actor_id, ads_service.AUDIT_ENTITY, audits)
            await ads_service.sync_initiatives(conn, [c["initiative_id"] for c in candidates], actor_id)

            # Campaigns this upload created, once nothing else is recorded against them
            campaign_ids = [UUID(i) for i in data.get("campaigns", [])]
            gone = await conn.fetch(
                """
                DELETE FROM initiatives i
                WHERE i.id = ANY($1::uuid[])
                  AND NOT EXISTS (SELECT 1 FROM ad_results r WHERE r.initiative_id = i.id)
                  AND NOT EXISTS (SELECT 1 FROM spend_lines s WHERE s.initiative_id = i.id)
                  AND NOT EXISTS (SELECT 1 FROM outcomes o WHERE o.initiative_id = i.id)
                RETURNING id, name, status, channel_id, sub_channel_id
                """,
                campaign_ids,
            )
            await ads_service.audit_many(
                conn, actor_id, INITIATIVE_AUDIT_ENTITY, [("DELETE", g["id"], dict(g), None) for g in gone]
            )

            await conn.execute(
                "UPDATE import_batches SET status = 'UNDONE', undone_at = now(), undone_by = $2 WHERE id = $1",
                batch_id, actor_id,
            )
            result = {
                "removed": len(removing), "restored": len(restoring),
                "skipped": len(created_ids) + len(previous) - len(candidates), "campaigns_removed": len(gone),
            }
            await log_audit(
                conn, actor_id, "UNDO_IMPORT", AUDIT_ENTITY, batch_id,
                old_value={"kind": KIND, "filename": batch["filename"], **result},
            )
    return {"batch_id": batch_id, "status": "UNDONE", **result}


# --- Template ---

async def build_template() -> bytes:
    """An .xlsx with Marketing's own header, and dropdowns for the brand, market, sub-channel and status."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    pool = get_pool()
    async with pool.acquire() as conn:
        brands = await conn.fetch("SELECT name FROM brands WHERE is_active ORDER BY sort_order, name")
        markets = await conn.fetch("SELECT name FROM markets WHERE is_active ORDER BY sort_order, name")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAMES[0]
    sheet.append(TEMPLATE_COLUMNS)
    bold = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
    sheet.freeze_panes = "A2"
    for index, header in enumerate(TEMPLATE_COLUMNS):
        column = sheet.cell(row=1, column=index + 1).column_letter
        sheet.column_dimensions[column].width = 34 if header in ("Campaign Name", "Report Link") else max(12, len(header) + 2)
    for column in ("C", "D", "U"):
        for row in sheet.iter_rows(min_row=2, max_row=1000, min_col=sheet[f"{column}1"].column, max_col=sheet[f"{column}1"].column):
            for cell in row:
                cell.number_format = "dd/mm/yyyy"

    lists = workbook.create_sheet("Lists")
    lists.append(["Brand", "Market/Region"])
    for cell in lists[1]:
        cell.font = bold
    for i in range(max(len(brands), len(markets))):
        lists.append([brands[i]["name"] if i < len(brands) else None, markets[i]["name"] if i < len(markets) else None])
    lists.column_dimensions["A"].width = 20
    lists.column_dimensions["B"].width = 20

    validations = {
        "A": f"=Lists!$A$2:$A${len(brands) + 1}",
        "B": f"=Lists!$B$2:$B${len(markets) + 1}",
        "E": '"Paid Search"',
        "F": '"Google Ads,Bing Ads"',
        "V": '"ongoing,Paused,Ended"',
    }
    for column, formula in validations.items():
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Choose a value from the list."
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}1000")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
