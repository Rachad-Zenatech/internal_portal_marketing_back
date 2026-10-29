"""Uploading an event's cost sheet (.xlsx or .csv) instead of retyping it.

Preview → confirm → (optionally) undo. The preview checks every row with the same rules as
POST /spend-lines (the shared checks in services/shared/spend_service.py) and keeps the rows that pass in an
import_batches row; confirming turns them into normal cost lines (source COST_SHEET_UPLOAD) in one statement.
Invoices are tracked by number only; no invoice files are stored.
"""
import io
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from uuid import UUID

import asyncpg

from models.spend_model import _MAX_AMOUNT
from postgresql_db.database import get_pool
from services.shared.ads_service import audit_many
from services.shared.import_parsing import (
    ROW_NUMBER,
    ImportFileError,
    InvalidValueError,
    match_columns,
    normalise_header,
    parse_amount,
    parse_currency,
    parse_date,
    read_table,
    uses_decimal_comma,
)
from services.shared.rbac_service import log_audit
from services.shared.spend_service import AUDIT_ENTITY as SPEND_AUDIT_ENTITY
from services.shared.spend_service import (
    InitiativeNotFoundError,
    SpendValidationError,
    check_category,
    check_paid_date,
    duplicate_invoice_message,
)

KIND = "COST_SHEET"
# Uploads this module confirms and undoes; ads and budget uploads have their own routes
_KINDS = ["COST_SHEET", "EVENTS"]
AUDIT_ENTITY = "ImportBatch"
MAX_FILE_BYTES = 5 * 1024 * 1024
PREVIEW_LIFETIME = timedelta(hours=24)

# Sheet header → our field (compared case-insensitively; spaces, underscores and dashes count the same).
# Set from the task brief: Marketing's real cost sheet hasn't been seen yet, so extend these when it is.
COLUMN_ALIASES: dict[str, list[str]] = {
    "description": ["Item", "Description", "Line item", "Details"],
    "category": ["Category"],
    "expense_type": ["Expense type", "Type"],
    "vendor": ["Vendor", "Supplier"],
    "estimate": ["Estimate", "Quote", "Budgeted", "Planned"],
    "actual": ["Actual", "Invoiced", "Final"],
    "status": ["Status"],
    # Card purchases come with a receipt rather than an invoice; either number goes here
    "invoice_ref": ["Invoice", "Invoice #", "Invoice No", "Invoice Number", "Receipt", "Receipt #", "Receipt No",
                    "Receipt Number", "Invoice or receipt #", "Invoice / receipt #", "Invoice/receipt #"],
    "date": ["Date", "Invoice date", "Receipt date", "Bill date", "Due date", "Expected date"],
    "paid_on": ["Paid On", "Paid date", "Date paid"],
    "payment_method": ["Payment Method"],
    "currency": ["Currency"],
    "cost_type": ["OPEX/CAPEX", "Cost type"],
    "note": ["Notes", "Note", "Comments"],
}
# The template's columns, in order
TEMPLATE_COLUMNS = [
    "Item", "Category", "Expense type", "Vendor", "Estimate", "Actual", "Status", "Invoice or receipt #",
    "Date", "Paid On", "Payment Method", "Currency", "OPEX/CAPEX", "Notes",
]

_STATUS_WORDS = {
    "estimate": "ESTIMATED", "estimated": "ESTIMATED", "quote": "ESTIMATED", "planned": "ESTIMATED",
    "committed": "COMMITTED", "ordered": "COMMITTED", "booked": "COMMITTED",
    "invoiced": "INVOICED",
    "paid": "PAID",
}
_STATUS_LABELS = {"ESTIMATED": "Estimate", "COMMITTED": "Committed", "INVOICED": "Invoiced", "PAID": "Paid"}
_PAYMENT_WORDS = {
    "ap invoice": "AP_INVOICE", "ap": "AP_INVOICE", "invoice": "AP_INVOICE", "supplier invoice": "AP_INVOICE",
    "corporate card": "CORPORATE_CARD", "card": "CORPORATE_CARD", "credit card": "CORPORATE_CARD",
    "other": "OTHER",
}
_PAYMENT_LABELS = {"AP_INVOICE": "AP invoice", "CORPORATE_CARD": "Corporate card", "OTHER": "Other"}


class ImportNotFoundError(Exception):
    """No upload with the requested id."""


class ImportForbiddenError(Exception):
    """Someone other than the uploader tried to confirm a preview."""


class ImportConflictError(Exception):
    """The upload isn't in a state that allows this; the message is safe to show to users."""


class ImportExpiredError(Exception):
    """The preview is too old to confirm."""


# Navan books flights, hotels and cars; its cost lines use Flights, Hotels, Ground transport (cars) and Other travel.
# When an event has Navan bookings, sheet rows in these categories are probably the same travel again (ground
# transport is left out: sheets list taxis and rideshares that Navan doesn't book).
NAVAN_OVERLAP = {"FLIGHTS": "flights", "HOTELS": "hotels", "OTHER_TRAVEL": "travel costs"}


# --- Checking one row ---

class _Context:
    """Everything a row check needs, loaded once per upload."""

    def __init__(self, initiative, categories, existing, columns: dict[str, str], decimal_comma: bool, navan: bool = False):
        self.currency = initiative["currency"]
        # The event already has travel from Navan
        self.navan = navan
        self.channel = initiative["channel"]
        self.start_date = initiative["start_date"]
        self.columns = columns
        self.decimal_comma = decimal_comma
        # Categories by normalised name and code, plus expense types that have exactly one category
        self.categories: dict[str, dict] = {}
        by_type: dict[str, list[dict]] = {}
        for c in categories:
            category = dict(c)
            self.categories.setdefault(normalise_header(category["name"]), category)
            self.categories.setdefault(normalise_header(category["code"]), category)
            by_type.setdefault(normalise_header(category["type_name"]), []).append(category)
        for type_name, members in by_type.items():
            if len(members) == 1:
                self.categories.setdefault(type_name, members[0])
        # Supplier invoices already recorded anywhere: (lower(vendor), invoice no.) → where
        self.existing = {(r["vendor_key"], r["invoice_ref"]): r for r in existing}
        self.seen_in_file: dict[tuple[str, str], int] = {}


def _cell(row: dict, ctx: _Context, field: str) -> str:
    header = ctx.columns.get(field)
    return str(row.get(header, "")).strip() if header else ""


def _check_row(row: dict, ctx: _Context) -> tuple[dict, dict | None]:
    """Check one sheet row. Returns (preview row, row to import or None)."""
    errors: list[str] = []
    warnings: list[str] = []
    row_number = row[ROW_NUMBER]

    item = _cell(row, ctx, "description")
    if not item:
        errors.append("Add the item (what the cost is for).")
    elif len(item) > 255:
        errors.append("Item is longer than 255 characters.")

    category = None
    raw_category = _cell(row, ctx, "category")
    if not raw_category:
        errors.append("Choose a category from the template's Categories sheet.")
    else:
        category = ctx.categories.get(normalise_header(raw_category))
        if category is None:
            errors.append(f"Unknown category '{raw_category}'. Choose one from the template's Categories sheet.")
    raw_type = _cell(row, ctx, "expense_type")
    if category:
        if raw_type and normalise_header(raw_type) not in (normalise_header(category["type_name"]), normalise_header(category["type_code"])):
            errors.append(f"{category['name']} is under {category['type_name']}, not {raw_type}")
        try:
            check_category(category, ctx.channel)
        except SpendValidationError as exc:
            errors.append(f"{category['name']}: {exc}")
        if ctx.navan and category["code"] in NAVAN_OVERLAP:
            warnings.append(
                f"Navan already has bookings for this event, so these {NAVAN_OVERLAP[category['code']]} may be counted twice. "
                "Leave Navan bookings out of the sheet."
            )

    def amount(field: str, label: str) -> Decimal | None:
        try:
            value = parse_amount(_cell(row, ctx, field), decimal_comma=ctx.decimal_comma)
        except InvalidValueError as exc:
            errors.append(f"{label}: {exc}")
            return None
        # The same limit as a cost typed in by hand
        if value is not None and abs(value) > _MAX_AMOUNT:
            errors.append(f"{label}: too large")
            return None
        return value

    estimate = amount("estimate", "Estimate")
    actual = amount("actual", "Actual")
    amounts_readable = not any(m.startswith(("Estimate:", "Actual:")) for m in errors)
    value = actual if actual is not None else estimate
    if value is None and amounts_readable:
        errors.append("Add an estimate or an actual amount")
    elif value == 0:
        errors.append("Amount can't be zero. Use a negative amount for a refund or credit.")

    def day(field: str, label: str) -> date | None:
        try:
            return parse_date(_cell(row, ctx, field))
        except InvalidValueError as exc:
            errors.append(f"{label}: {exc}")
            return None

    incurred_on = day("date", "Date")
    paid_on = day("paid_on", "Paid On")

    raw_status = _cell(row, ctx, "status")
    if raw_status:
        status = _STATUS_WORDS.get(normalise_header(raw_status))
        if status is None:
            errors.append(f"Unknown status '{raw_status}'. Use Estimate, Committed, Invoiced or Paid.")
    else:
        # Blank status: worked out from what's filled in
        status = "PAID" if paid_on else "INVOICED" if actual is not None else "ESTIMATED"
    if status in ("INVOICED", "PAID") and actual is None and estimate is not None:
        warnings.append("No actual amount, so the estimate was used.")
    if status:
        try:
            check_paid_date(status, paid_on)
        except SpendValidationError:
            if not any(m.startswith("Paid On:") for m in errors):
                errors.append("Paid needs a Paid On date.")
    if paid_on and status and status != "PAID":
        warnings.append(f"Paid On ignored because the status is {_STATUS_LABELS[status]}")

    if incurred_on is None and not _cell(row, ctx, "date"):
        if ctx.start_date:
            incurred_on = ctx.start_date
            warnings.append("No date, so the event's start date was used.")
        else:
            incurred_on = datetime.now(timezone.utc).date()
            warnings.append("No date, so today's date was used.")

    try:
        currency = parse_currency(_cell(row, ctx, "currency"), ctx.currency)
    except InvalidValueError as exc:
        errors.append(f"Currency: {exc}")
        currency = None

    raw_method = _cell(row, ctx, "payment_method")
    payment_method = _PAYMENT_WORDS.get(normalise_header(raw_method)) if raw_method else None
    if raw_method and payment_method is None:
        errors.append(f"Unknown payment method '{raw_method}'. Use AP invoice, Corporate card or Other.")

    raw_cost_type = _cell(row, ctx, "cost_type").upper()
    if raw_cost_type and raw_cost_type not in ("OPEX", "CAPEX"):
        errors.append(f"OPEX/CAPEX must be OPEX or CAPEX, not '{_cell(row, ctx, 'cost_type')}'")
    cost_type = raw_cost_type or (category["default_cost_type"] if category else None)

    vendor = _cell(row, ctx, "vendor") or None
    if vendor and len(vendor) > 255:
        errors.append("Vendor is longer than 255 characters.")
    invoice_ref = _cell(row, ctx, "invoice_ref") or None
    if invoice_ref and len(invoice_ref) > 100:
        errors.append("Invoice # is longer than 100 characters.")
    note = _cell(row, ctx, "note") or None

    # The same supplier invoice recorded already, or earlier in this file
    duplicate = None
    key = (vendor.lower(), invoice_ref) if vendor and invoice_ref else None
    if key and key in ctx.existing:
        other = ctx.existing[key]
        duplicate = duplicate_invoice_message(invoice_ref, other["vendor"], other["name"], other["event_code"])
    elif key and key in ctx.seen_in_file:
        duplicate = f"Invoice {invoice_ref} from {vendor} appears twice in this file (row {ctx.seen_in_file[key]})."

    if errors:
        row_status, messages = "INVALID", errors + ([duplicate] if duplicate else [])
    elif duplicate:
        row_status, messages = "DUPLICATE", [duplicate]
    else:
        row_status, messages = ("READY_WITH_WARNINGS" if warnings else "READY"), warnings
        if key:
            ctx.seen_in_file[key] = row_number

    preview = {
        "row_number": row_number,
        "status": row_status,
        "messages": messages,
        "item": item or None,
        "category": category["name"] if category else raw_category or None,
        "expense_type": category["type_name"] if category else raw_type or None,
        "vendor": vendor,
        "estimate": float(estimate) if estimate is not None else None,
        "actual": float(actual) if actual is not None else None,
        "status_value": _STATUS_LABELS.get(status) if status else raw_status or None,
        "invoice_ref": invoice_ref,
        "date": incurred_on,
        "paid_on": paid_on,
        "currency": currency,
    }
    if row_status not in ("READY", "READY_WITH_WARNINGS"):
        return preview, None
    staged = {
        "row_number": row_number,
        "category_id": category["id"],
        "description": item,
        "vendor": vendor,
        "note": note,
        "status": status,
        "estimated_amount": str(estimate) if estimate is not None else None,
        "amount": str(value),
        "currency": currency,
        "incurred_on": incurred_on.isoformat(),
        "invoice_ref": invoice_ref,
        "paid_on": paid_on.isoformat() if paid_on and status == "PAID" else None,
        "payment_method": payment_method,
        "cost_type": cost_type,
    }
    return preview, staged


# --- The upload a new one replaces ---

# The event's current cost sheet upload, with its version and what's on the event from it now
_CURRENT_UPLOAD_SQL = f"""
    SELECT b.id, b.filename, b.confirmed_at,
           (SELECT count(*) FROM import_batches x
            WHERE x.initiative_id = b.initiative_id AND x.kind = b.kind AND x.status IN ('CONFIRMED', 'REPLACED')
              AND x.confirmed_at <= b.confirmed_at) AS version,
           count(s.id) AS line_count,
           count(s.id) FILTER (WHERE s.updated_at > s.created_at + interval '1 second') AS edited_count,
           count(s.id) FILTER (WHERE s.reconciled_month IS NOT NULL) AS reconciled_count
    FROM import_batches b
    LEFT JOIN spend_lines s ON s.import_batch_id = b.id
    WHERE b.initiative_id = $1 AND b.kind = '{KIND}' AND b.status = 'CONFIRMED'
    GROUP BY b.id
    ORDER BY b.confirmed_at DESC
    LIMIT 1
"""


# --- Preview ---

_CATEGORIES_SQL = """
    SELECT sc.id, sc.code, sc.name, sc.manual_entry_allowed, sc.applies_to, sc.default_cost_type,
           et.code AS type_code, et.name AS type_name, et.sort_order AS type_order, sc.sort_order
    FROM spend_categories sc
    JOIN expense_types et ON et.id = sc.expense_type_id
    WHERE sc.is_active AND et.is_active
    ORDER BY et.sort_order, sc.sort_order
"""


async def preview_cost_sheet(initiative_id: UUID, filename: str, data: bytes, actor_id: UUID) -> dict:
    """Check every row of an uploaded cost sheet and keep the rows that pass for confirming later."""
    known = [name for names in COLUMN_ALIASES.values() for name in names]
    rows = read_table(filename, data, known)
    if not rows:
        raise ImportFileError("The sheet has a header but no cost rows.")
    headers = [h for h in rows[0] if h != ROW_NUMBER]
    match = match_columns(headers, COLUMN_ALIASES)
    missing = [label for field, label in (("description", "Item"), ("category", "Category")) if field not in match.fields]
    if missing:
        raise ImportFileError(f"The sheet needs {' and '.join(missing)} column{'s' if len(missing) > 1 else ''}. Start from the template.")
    if "estimate" not in match.fields and "actual" not in match.fields:
        raise ImportFileError("The sheet needs an Estimate or an Actual column. Start from the template.")
    amount_cells = [row.get(match.fields[f], "") for row in rows for f in ("estimate", "actual") if f in match.fields]

    vendor_keys = sorted({str(row.get(match.fields["vendor"], "")).strip().lower() for row in rows} - {""}) if "vendor" in match.fields else []
    pool = get_pool()
    async with pool.acquire() as conn:
        # A fixed number of queries, however long the sheet is
        initiative = await conn.fetchrow(
            "SELECT i.currency, i.start_date, c.code AS channel FROM initiatives i JOIN channels c ON c.id = i.channel_id WHERE i.id = $1",
            initiative_id,
        )
        if initiative is None:
            raise InitiativeNotFoundError
        categories = await conn.fetch(_CATEGORIES_SQL)
        current = await conn.fetchrow(_CURRENT_UPLOAD_SQL, initiative_id)
        existing = await conn.fetch(
            """
            SELECT lower(s.vendor) AS vendor_key, s.vendor, s.invoice_ref, i.name, ed.event_code
            FROM spend_lines s
            JOIN initiatives i ON i.id = s.initiative_id
            LEFT JOIN event_details ed ON ed.initiative_id = i.id
            WHERE lower(s.vendor) = ANY($1::text[]) AND s.invoice_ref IS NOT NULL
              -- The current upload's costs are replaced on confirm, so its invoices aren't duplicates (reconciled ones stay)
              AND NOT COALESCE(s.import_batch_id = $2::uuid AND s.reconciled_month IS NULL, false)
            """,
            vendor_keys, current["id"] if current else None,
        )
        navan = await conn.fetchval(
            "SELECT EXISTS (SELECT 1 FROM spend_lines WHERE initiative_id = $1 AND source = 'NAVAN')", initiative_id
        )

        ctx = _Context(initiative, categories, existing, match.fields, uses_decimal_comma(amount_cells), navan=navan)
        previews, staged = [], []
        for row in rows:
            preview, ready = _check_row(row, ctx)
            previews.append(preview)
            if ready:
                staged.append(ready)

        counts = {
            "total": len(previews),
            "ready": len(staged),
            "warnings": sum(p["status"] == "READY_WITH_WARNINGS" for p in previews),
            "duplicates": sum(p["status"] == "DUPLICATE" for p in previews),
            "invalid": sum(p["status"] == "INVALID" for p in previews),
        }
        batch_id = await conn.fetchval(
            """
            INSERT INTO import_batches (kind, initiative_id, filename, status, total_rows, ready_rows, error_rows,
                                        duplicate_rows, staged_rows, created_by)
            VALUES ($1, $2, $3, 'PREVIEW', $4, $5, $6, $7, $8::jsonb, $9)
            RETURNING id
            """,
            KIND, initiative_id, (filename or "upload")[:255], counts["total"], counts["ready"], counts["invalid"],
            counts["duplicates"], json.dumps(staged), actor_id,
        )

    # What confirming would add, in the event's currency (other currencies are never converted)
    same_currency = [s for s in staged if s["currency"] == ctx.currency]
    total = lambda statuses: float(sum((Decimal(s["amount"]) for s in same_currency if s["status"] in statuses), Decimal(0)))  # noqa: E731
    return {
        "batch_id": batch_id,
        "filename": filename,
        "counts": counts,
        "ignored_columns": match.ignored,
        "rows": previews,
        "totals_preview": {
            "expected_total": total({"ESTIMATED", "COMMITTED", "INVOICED", "PAID"}),
            "committed": total({"COMMITTED", "INVOICED", "PAID"}),
            "actual": total({"INVOICED", "PAID"}),
        },
        "replaces": (
            {
                "batch_id": current["id"], "filename": current["filename"], "confirmed_at": current["confirmed_at"],
                "version": current["version"], "line_count": current["line_count"],
                "edited_count": current["edited_count"], "reconciled_count": current["reconciled_count"],
            }
            if current else None
        ),
    }


# --- Confirm and undo ---

def _is_expired(created_at: datetime) -> bool:
    return created_at < datetime.now(timezone.utc) - PREVIEW_LIFETIME


def _staged(batch) -> list[dict]:
    raw = batch["staged_rows"]
    return json.loads(raw) if isinstance(raw, str) else raw


async def _insert_lines(conn, initiative_id: UUID, batch_id: UUID, staged: list[dict], actor_id: UUID) -> list:
    """Add a batch's checked rows as cost lines in one statement, and audit each like a cost added by hand.
    Rows whose supplier invoice is already recorded are skipped."""
    column = lambda key, convert=None: [  # noqa: E731
        (convert(r[key]) if convert and r[key] is not None else r[key]) for r in staged
    ]
    try:
        inserted = await conn.fetch(
            """
            INSERT INTO spend_lines (initiative_id, category_id, description, vendor, note, status,
                                     estimated_amount, amount, currency, incurred_on, invoice_ref, paid_on,
                                     payment_method, cost_type, source, import_batch_id, created_by, updated_by)
            SELECT $1, r.category_id, r.description, r.vendor, r.note, r.status, r.estimated_amount, r.amount,
                   r.currency, r.incurred_on, r.invoice_ref, r.paid_on, r.payment_method, r.cost_type,
                   'COST_SHEET_UPLOAD', $2, $3, $3
            FROM unnest($4::int[], $5::text[], $6::text[], $7::text[], $8::text[], $9::numeric[], $10::numeric[],
                        $11::text[], $12::date[], $13::text[], $14::date[], $15::text[], $16::text[])
                 AS r(category_id, description, vendor, note, status, estimated_amount, amount, currency,
                      incurred_on, invoice_ref, paid_on, payment_method, cost_type)
            -- Someone may have recorded the same supplier invoice since the preview: skip those rows
            ON CONFLICT (lower(vendor), invoice_ref) WHERE vendor IS NOT NULL AND invoice_ref IS NOT NULL DO NOTHING
            RETURNING *
            """,
            initiative_id, batch_id, actor_id,
            column("category_id"), column("description"), column("vendor"), column("note"), column("status"),
            column("estimated_amount", Decimal), column("amount", Decimal), column("currency"),
            column("incurred_on", date.fromisoformat), column("invoice_ref"), column("paid_on", date.fromisoformat),
            column("payment_method"), column("cost_type"),
        )
    except (asyncpg.CheckViolationError, asyncpg.ForeignKeyViolationError, asyncpg.DataError) as exc:
        raise ImportConflictError("Some rows no longer pass the checks. Upload the file again.") from exc
    await audit_many(conn, actor_id, SPEND_AUDIT_ENTITY, [("CREATE", line["id"], None, dict(line)) for line in inserted])
    return inserted


async def confirm_import(batch_id: UUID, actor_id: UUID) -> dict:
    """Turn a preview's rows into cost lines in one statement, then mark the upload confirmed."""
    expired = False
    pool = get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            batch = await conn.fetchrow(
                "SELECT kind, initiative_id, filename, status, staged_rows, created_by, created_at FROM import_batches "
                "WHERE id = $1 AND kind = ANY($2::text[]) FOR UPDATE",
                batch_id, _KINDS,
            )
            if batch is None:
                raise ImportNotFoundError
            if batch["created_by"] != actor_id:
                raise ImportForbiddenError
            if batch["status"] != "PREVIEW":
                raise ImportConflictError(f"This upload is already {batch['status'].lower()}.")
            if _is_expired(batch["created_at"]):
                await conn.execute("UPDATE import_batches SET status = 'EXPIRED' WHERE id = $1", batch_id)
                expired = True
            else:
                staged = _staged(batch)
                # A new cost sheet for the event replaces its current one: people keep one living Excel file and upload
                # it again. The current upload's costs go (reconciled ones stay, as Finance has signed them off).
                replaced_rows, replaced_id = 0, None
                if batch["kind"] == KIND:
                    replaced_id = await conn.fetchval(
                        f"""
                        SELECT id FROM import_batches
                        WHERE initiative_id = $1 AND kind = '{KIND}' AND status = 'CONFIRMED'
                        ORDER BY confirmed_at DESC LIMIT 1 FOR UPDATE
                        """,
                        batch["initiative_id"],
                    )
                if replaced_id is not None:
                    gone = await conn.fetch(
                        "DELETE FROM spend_lines WHERE import_batch_id = $1 AND reconciled_month IS NULL RETURNING *", replaced_id
                    )
                    replaced_rows = len(gone)
                    await audit_many(conn, actor_id, SPEND_AUDIT_ENTITY, [("DELETE", line["id"], dict(line), None) for line in gone])
                    await conn.execute(
                        "UPDATE import_batches SET status = 'REPLACED', undone_at = now(), undone_by = $2 WHERE id = $1",
                        replaced_id, actor_id,
                    )
                    await conn.execute(
                        "UPDATE import_batches SET undo_data = $2::jsonb WHERE id = $1",
                        batch_id, json.dumps({"replaced_batch_id": str(replaced_id), "replaced_rows": replaced_rows}),
                    )
                inserted = await _insert_lines(conn, batch["initiative_id"], batch_id, staged, actor_id)
                imported, skipped = len(inserted), len(staged) - len(inserted)
                # Linking the new lines to monthly budgets comes with the budget tracker (Part 4)
                await conn.execute(
                    """
                    UPDATE import_batches
                    SET status = 'CONFIRMED', imported_rows = $2, confirmed_at = now(), confirmed_by = $3
                    WHERE id = $1
                    """,
                    batch_id, imported, actor_id,
                )
                await log_audit(
                    conn, actor_id, "IMPORT", AUDIT_ENTITY, batch_id,
                    new_value={"kind": batch["kind"], "filename": batch["filename"], "initiative_id": batch["initiative_id"],
                               "imported_rows": imported, "skipped": skipped, "replaced_batch_id": replaced_id,
                               "replaced_rows": replaced_rows},
                )
    if expired:
        raise ImportExpiredError
    return {"batch_id": batch_id, "status": "CONFIRMED", "imported_rows": imported, "skipped": skipped, "replaced_rows": replaced_rows}


async def undo_import(batch_id: UUID, actor_id: UUID) -> dict:
    """Remove every cost line an upload created (including ones edited since), unless Finance reconciled any.
    Undoing an upload that replaced the event's previous one puts the previous upload's costs back.

    Only the uploader can confirm their own preview, but anyone who can edit costs can undo a confirmed upload,
    so a bad sheet can be pulled while its uploader is away. Each removed line's last values are audited.
    """
    pool = get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            batch = await conn.fetchrow(
                "SELECT kind, initiative_id, filename, status, undo_data FROM import_batches WHERE id = $1 AND kind = ANY($2::text[]) FOR UPDATE",
                batch_id, _KINDS,
            )
            if batch is None:
                raise ImportNotFoundError
            if batch["status"] != "CONFIRMED":
                raise ImportConflictError("Only a confirmed upload can be undone.")
            reconciled = await conn.fetchval(
                "SELECT count(*) FROM spend_lines WHERE import_batch_id = $1 AND reconciled_month IS NOT NULL", batch_id
            )
            if reconciled:
                verb = "is" if reconciled == 1 else "are"
                raise ImportConflictError(f"{reconciled} of these costs {verb} already reconciled. Ask Finance to reopen them first.")
            gone = await conn.fetch("DELETE FROM spend_lines WHERE import_batch_id = $1 RETURNING *", batch_id)
            removed = len(gone)
            await audit_many(conn, actor_id, SPEND_AUDIT_ENTITY, [("DELETE", line["id"], dict(line), None) for line in gone])
            await conn.execute(
                "UPDATE import_batches SET status = 'UNDONE', undone_at = now(), undone_by = $2 WHERE id = $1", batch_id, actor_id
            )
            # Put back the upload this one replaced, from its checked rows (reconciled costs it kept aren't added twice)
            undo_data = batch["undo_data"]
            undo_data = json.loads(undo_data) if isinstance(undo_data, str) else (undo_data or {})
            restored = 0
            if undo_data.get("replaced_batch_id"):
                previous = await conn.fetchrow(
                    "SELECT id, initiative_id, staged_rows FROM import_batches WHERE id = $1 AND status = 'REPLACED' FOR UPDATE",
                    UUID(undo_data["replaced_batch_id"]),
                )
                if previous is not None:
                    kept = await conn.fetchval("SELECT count(*) FROM spend_lines WHERE import_batch_id = $1", previous["id"])
                    rows = _staged(previous)
                    if kept:
                        # Its reconciled costs never left; only rows with a supplier invoice can be told apart from them,
                        # so the rest come back only when nothing was kept
                        rows = [r for r in rows if r.get("vendor") and r.get("invoice_ref")]
                    restored = len(await _insert_lines(conn, previous["initiative_id"], previous["id"], rows, actor_id))
                    await conn.execute(
                        "UPDATE import_batches SET status = 'CONFIRMED', undone_at = NULL, undone_by = NULL WHERE id = $1", previous["id"]
                    )
            await log_audit(
                conn, actor_id, "UNDO_IMPORT", AUDIT_ENTITY, batch_id,
                old_value={"kind": batch["kind"], "filename": batch["filename"], "initiative_id": batch["initiative_id"],
                           "removed_rows": removed, "restored_rows": restored},
            )
    return {"batch_id": batch_id, "status": "UNDONE", "removed_rows": removed, "restored_rows": restored}


# --- History ---

async def list_imports(initiative_id: UUID, limit: int = 50) -> list[dict]:
    """An initiative's recent uploads, newest first, with how many of their lines remain and were edited."""
    pool = get_pool()
    async with pool.acquire() as conn:
        if not await conn.fetchval("SELECT EXISTS (SELECT 1 FROM initiatives WHERE id = $1)", initiative_id):
            raise InitiativeNotFoundError
        rows = await conn.fetch(
            """
            SELECT b.id, b.kind, b.filename,
                   CASE WHEN b.status = 'PREVIEW' AND b.created_at < now() - interval '24 hours' THEN 'EXPIRED' ELSE b.status END AS status,
                   b.total_rows, b.ready_rows, b.error_rows, b.duplicate_rows, b.imported_rows,
                   cu.full_name AS created_by_name, b.created_at,
                   b.confirmed_at, fu.full_name AS confirmed_by_name,
                   b.undone_at, uu.full_name AS undone_by_name,
                   COALESCE(l.line_count, 0) AS line_count, COALESCE(l.edited_count, 0) AS edited_count,
                   -- v1, v2… among the event's confirmed uploads (including replaced ones)
                   CASE WHEN b.status IN ('CONFIRMED', 'REPLACED') THEN
                       (SELECT count(*) FROM import_batches x
                        WHERE x.initiative_id = b.initiative_id AND x.kind = b.kind AND x.status IN ('CONFIRMED', 'REPLACED')
                          AND x.confirmed_at <= b.confirmed_at)::int
                   END AS version
            FROM import_batches b
            LEFT JOIN users cu ON cu.id = b.created_by
            LEFT JOIN users fu ON fu.id = b.confirmed_by
            LEFT JOIN users uu ON uu.id = b.undone_by
            LEFT JOIN LATERAL (
                SELECT count(*) AS line_count,
                       count(*) FILTER (WHERE s.updated_at > s.created_at + interval '1 second') AS edited_count
                FROM spend_lines s WHERE s.import_batch_id = b.id
            ) l ON true
            WHERE b.initiative_id = $1 AND b.kind = $2
            ORDER BY b.created_at DESC
            LIMIT $3
            """,
            initiative_id, KIND, limit,
        )
    return [dict(r) for r in rows]


# --- Template ---

async def build_template() -> bytes:
    """An .xlsx to start from: the columns in order, an example booth, and the categories to choose from."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    pool = get_pool()
    async with pool.acquire() as conn:
        categories = await conn.fetch(_CATEGORIES_SQL)
    # Categories people can type in for an event: everything but ads-only media (Navan adds its own bookings)
    choices = [c for c in categories if c["manual_entry_allowed"] and c["applies_to"] in ("ALL", "EVENTS")]

    workbook = Workbook()
    costs = workbook.active
    costs.title = "Costs"
    costs.append(TEMPLATE_COLUMNS)
    bold = Font(bold=True)
    for cell in costs[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
    costs.freeze_panes = "A2"

    today = date.today()
    show = today + timedelta(days=60)
    examples = [
        ("10x20 booth space", "Booth space", "Space & sponsorship", "Show organiser", 8000, 8000, "Paid", "EX-INV-001",
         today, today + timedelta(days=14), "AP invoice", "", "", "Deposit paid when booking"),
        ("Modular booth build", "Booth build or rental", "Booth / exhibit", "ExhibitWorks", 4000, 4500, "Invoiced", "EX-INV-002",
         today + timedelta(days=30), None, "", "", "", ""),
        ("10x20 grey carpet", "Carpet & flooring", "Show services", "Expo Show Services", 600, None, "Committed", "",
         show, None, "", "", "", ""),
        ("Booth lighting package", "Lighting", "Show services", "Expo Show Services", 450, None, "Committed", "",
         show, None, "", "", "", ""),
        ("Electrical, 2 outlets", "Electrical", "Show services", "Expo Show Services", 350, None, "Committed", "",
         show, None, "", "", "", ""),
        ("Exhibitor badges x4", "Badges & passes", "Registration & hospitality", "Show organiser", 300, 300, "Paid", "",
         today, today, "Corporate card", "", "", ""),
        ("Booth freight", "Freight", "Shipping & logistics", "Expo Freight Co.", 1800, None, "Estimate", "",
         show - timedelta(days=7), None, "", "", "", "Quote only"),
        ("Banners and brochures", "Printing & signage", "Marketing & promotion", "PrintHouse", 900, None, "Committed", "",
         show - timedelta(days=14), None, "", "", "", ""),
    ]
    for example in examples:
        costs.append(list(example))
    for row in costs.iter_rows(min_row=2, min_col=9, max_col=10):
        for cell in row:
            cell.number_format = "dd/mm/yyyy"
    for column, width in zip("ABCDEFGHIJKLMN", (28, 24, 24, 22, 12, 12, 12, 14, 13, 13, 16, 10, 12, 30)):
        costs.column_dimensions[column].width = width

    sheet = workbook.create_sheet("Categories")
    sheet.append(["Expense type", "Category"])
    for cell in sheet[1]:
        cell.font = bold
    for c in choices:
        sheet.append([c["type_name"], c["name"]])
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 28

    # Dropdowns on the Costs sheet (rows 2–1000)
    last = len(choices) + 1
    validations = {
        "B": f"=Categories!$B$2:$B${last}",
        "G": '"Estimate,Committed,Invoiced,Paid"',
        "K": '"AP invoice,Corporate card,Other"',
        "M": '"OPEX,CAPEX"',
    }
    for column, formula in validations.items():
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Choose a value from the list."
        costs.add_data_validation(validation)
        validation.add(f"{column}2:{column}1000")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
