"""Uploading Marketing's budget tracker sheet ("Budget Tracker") instead of adding each budget by hand.

Preview → confirm → (optionally) undo, like the PPC sheet upload (services/ads_import_service.py).
Each row is a budget for a month, brand and market (and optionally a channel). The preview says whether it
would add a budget, update one, or change nothing. Confirming adds and updates the budgets and adds the row's
note as a comment on the budget. Undo removes the budgets the upload added and gives the ones it changed their
previous values back.

Only the approved budget, status and notes are read. Spend, carryover and unspent are worked out from costs
and money moves, so those columns are ignored. The sheet's Spend Progress gives the status: Canceled and Not
committed carry over; Complete, In Progress and Not Started are all approved budgets whose progress the portal
works out from spend.
"""
import calendar
import json
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from uuid import UUID

from postgresql_db.database import get_pool
from services.budget_service import BUDGET_ENTITY, COMMENT_ENTITY, TBD_ALL_BRAND
from services.cost_sheet_import_service import (
    ImportConflictError,
    ImportExpiredError,
    ImportForbiddenError,
    ImportNotFoundError,
)
from services.import_parsing import (
    ROW_NUMBER,
    ImportFileError,
    InvalidValueError,
    match_columns,
    normalise_header,
    parse_amount,
    parse_currency,
    parse_date,
    read_table,
    uses_decimal_comma,
)
from services.rbac_service import log_audit

KIND = "BUDGETS"
AUDIT_ENTITY = "ImportBatch"
MAX_FILE_BYTES = 5 * 1024 * 1024
PREVIEW_LIFETIME = timedelta(hours=24)
SHEET_NAMES = ["Budget Tracker", "Budget tracker", "Budgets"]
DEFAULT_CURRENCY = "USD"

# Sheet header → our field (compared case-insensitively; spaces, underscores and dashes count the same)
COLUMN_ALIASES: dict[str, list[str]] = {
    "month": ["Month"],
    "brand": ["Brand"],
    "market": ["Market / Region", "Market/Region", "Market", "Region"],
    "channel": ["Channel"],
    "amount": ["Approved Budget", "Approved Budget ($)", "Budget", "Amount"],
    "currency": ["Currency"],
    "progress": ["Spend Progress", "Status", "Progress"],
    "note": ["Notes", "Note", "Comments"],
}
# Marketing's own header, in order, so they can keep their layout
TEMPLATE_COLUMNS = [
    "Month", "Brand", "Market / Region", "Approved Budget", "Spend Progress", "Actual Spend to Date",
    "Carryover In\n(from prior month)", "Unspent\n(carries to next month)", "Carryover Spent\nThis Month", "Notes",
    "Last Update",
]

# Spend Progress → budget status. Progress itself is worked out from spend, so the rest are all approved.
_PROGRESS_WORDS = {
    "canceled": "CANCELLED", "cancelled": "CANCELLED",
    "not committed": "NOT_COMMITTED", "uncommitted": "NOT_COMMITTED", "pending": "NOT_COMMITTED",
    "complete": "APPROVED", "completed": "APPROVED", "in progress": "APPROVED", "not started": "APPROVED",
    "approved": "APPROVED", "ongoing": "APPROVED",
}
_PROGRESS_LABELS = "Complete, In Progress, Not Started, Not committed or Canceled"
_BLANK_CELLS = {"", ".", "-", "–", "—", "n/a", "#n/a"}
_ALL_CHANNELS = {"all", "all channels", "any"}
_MONTH_NUMBERS = {
    **{name.lower(): i for i, name in enumerate(calendar.month_name) if name},
    **{name.lower(): i for i, name in enumerate(calendar.month_abbr) if name},
    "sept": 9,
}
_MAX_AMOUNT = Decimal("999999999999")
_MAX_NOTE = 2000
_VALUE_COLUMNS = ("amount", "currency", "status")


def _join_words(words: list[str]) -> str:
    return words[0] if len(words) == 1 else f"{', '.join(words[:-1])} and {words[-1]}"


def _is_blank(text: str) -> bool:
    return text.strip().lower() in _BLANK_CELLS


def parse_month(text: str, year: int) -> date:
    """'Aug', 'August', 'Sep 2026', 'Sep-26', '2026-09', '09/2026' or a date cell → the first of that month.

    A month without a year uses `year`. Anything else raises InvalidValueError.
    """
    text = text.strip()
    named = re.fullmatch(r"([A-Za-z]+)\.?(?:[\s\-/',]+(\d{4}|\d{2}))?", text)
    if named:
        number = _MONTH_NUMBERS.get(named.group(1).lower())
        if number is None:
            raise InvalidValueError(f"'{text}' isn't a month")
        given = named.group(2)
        return date(int(given) + (2000 if len(given) == 2 else 0) if given else year, number, 1)
    numeric = re.fullmatch(r"(\d{4})-(\d{1,2})|(\d{1,2})/(\d{4})", text)
    if numeric:
        y, m = (numeric.group(1), numeric.group(2)) if numeric.group(1) else (numeric.group(4), numeric.group(3))
        if not 1 <= int(m) <= 12:
            raise InvalidValueError(f"'{text}' isn't a month")
        return date(int(y), int(m), 1)
    day = parse_date(text)
    if day is None:
        raise InvalidValueError("Add the month")
    return day.replace(day=1)


# --- Checking one row ---

class _Context:
    """Everything a row check needs, loaded once per upload."""

    def __init__(self, columns: dict[str, str], decimal_comma: bool, year: int, currency: str, brands, markets, channels):
        self.columns = columns
        self.decimal_comma = decimal_comma
        self.year = year
        self.currency = currency
        self.brands = _lookup(brands)
        tbd = next((dict(b) for b in brands if b["code"] == TBD_ALL_BRAND), None)
        for word in ("tbd", "all", "tbd all", "tbd/all"):
            if tbd:
                self.brands.setdefault(word, tbd)
        self.markets = _lookup(markets)
        self.channels = _lookup(channels)


def _lookup(rows) -> dict[str, dict]:
    found: dict[str, dict] = {}
    for r in rows:
        found.setdefault(normalise_header(r["name"]), dict(r))
        found.setdefault(normalise_header(r["code"]), dict(r))
    return found


def _cell(row: dict, ctx: _Context, field: str) -> str:
    header = ctx.columns.get(field)
    return str(row.get(header, "")).strip() if header else ""


def _currency_in(text: str) -> str | None:
    """The currency an amount is written in, e.g. '£7,500' → GBP, 'A$ 50' → AUD; None for a bare number."""
    found = re.match(r"\s*((?:[A-Za-z]{1,2})?[$£€₩])", text) or re.search(r"(?<=[\d\s])([A-Za-z]{3})\s*$", text)
    if not found:
        return None
    try:
        return parse_currency(found.group(1), DEFAULT_CURRENCY)
    except InvalidValueError:
        return None


def _check_row(row: dict, ctx: _Context) -> dict:
    """Read and check one sheet row. Returns the preview row plus what would be imported (under '_staged')."""
    errors: list[str] = []
    warnings: list[str] = []

    month = None
    raw_month = _cell(row, ctx, "month")
    if _is_blank(raw_month):
        errors.append("Add the month.")
    else:
        try:
            month = parse_month(raw_month, ctx.year)
        except (InvalidValueError, ValueError):
            errors.append(f"'{raw_month}' isn't a month. Use Aug, August or Aug 2026.")

    raw_brand = _cell(row, ctx, "brand")
    brand = ctx.brands.get(normalise_header(raw_brand)) if raw_brand else None
    if _is_blank(raw_brand):
        errors.append("Add the brand, or TBD/All.")
    elif brand is None:
        errors.append(f"Unknown brand '{raw_brand}'.")
    raw_market = _cell(row, ctx, "market")
    market = ctx.markets.get(normalise_header(raw_market)) if raw_market else None
    if _is_blank(raw_market):
        errors.append("Add the market or region.")
    elif market is None:
        errors.append(f"Unknown market '{raw_market}'.")
    raw_channel = _cell(row, ctx, "channel")
    if _is_blank(raw_channel) or normalise_header(raw_channel) in _ALL_CHANNELS:
        raw_channel = ""
    channel = None
    if raw_channel:
        channel = ctx.channels.get(normalise_header(raw_channel))
        if channel is None:
            errors.append(f"Unknown channel '{raw_channel}'.")

    raw_progress = _cell(row, ctx, "progress")
    status = "APPROVED"
    if not _is_blank(raw_progress):
        status = _PROGRESS_WORDS.get(normalise_header(raw_progress))
        if status is None:
            errors.append(f"Unknown spend progress '{raw_progress}'. Use {_PROGRESS_LABELS}.")

    raw_amount = _cell(row, ctx, "amount")
    amount = None
    if _is_blank(raw_amount):
        amount = Decimal("0.00")
        if status != "CANCELLED":
            warnings.append("No approved budget in the sheet, so it's added as 0.")
    else:
        try:
            amount = parse_amount(raw_amount, decimal_comma=ctx.decimal_comma)
        except InvalidValueError as exc:
            errors.append(f"Approved Budget: {exc}")
        if amount is not None:
            if amount < 0:
                errors.append("Approved Budget can't be negative.")
                amount = None
            elif amount > _MAX_AMOUNT:
                errors.append("Approved Budget is too large.")
                amount = None
            else:
                amount = amount.quantize(Decimal("0.01"))

    currency = ctx.currency
    raw_currency = _cell(row, ctx, "currency")
    if not _is_blank(raw_currency):
        try:
            currency = parse_currency(raw_currency, ctx.currency)
        except InvalidValueError as exc:
            errors.append(f"Currency: {exc}")
    else:
        currency = _currency_in(raw_amount) or ctx.currency

    note = " ".join(_cell(row, ctx, "note").split())
    if _is_blank(note):
        note = ""
    if len(note) > _MAX_NOTE:
        warnings.append(f"The note is longer than {_MAX_NOTE} characters, so it was cut short.")
        note = note[:_MAX_NOTE]

    is_tbd = brand is not None and brand["code"] == TBD_ALL_BRAND
    preview = {
        "row_number": row[ROW_NUMBER],
        "status": "INVALID" if errors else "NEW",
        "messages": errors + warnings,
        "month": month,
        "brand": "TBD/All" if is_tbd else (brand["name"] if brand else raw_brand or None),
        "market": market["name"] if market else raw_market or None,
        "channel": channel["name"] if channel else raw_channel or None,
        "amount": float(amount) if amount is not None else None,
        "currency": currency,
        "budget_status": status,
        "sheet_progress": raw_progress or None,
        "note": note or None,
        "budget_id": None,
        "_staged": None,
    }
    if not errors:
        preview["_staged"] = {
            "row_number": row[ROW_NUMBER],
            "month": month,
            "brand_id": None if is_tbd else brand["id"],
            "market_id": market["id"],
            "channel_id": channel["id"] if channel else None,
            "amount": amount,
            "currency": currency,
            "status": status,
            "note": note or None,
        }
    return preview


def _key(row: dict) -> tuple:
    return row["month"], row["brand_id"] or 0, row["market_id"], row["channel_id"] or 0


def _changed(current: dict, staged: dict) -> list[str]:
    return [c for c in _VALUE_COLUMNS if (Decimal(current[c]) != Decimal(staged[c]) if c == "amount" else current[c] != staged[c])]


def _json_ready(staged: dict) -> dict:
    return {k: (v.isoformat() if isinstance(v, date) else str(v) if isinstance(v, (Decimal, UUID)) else v) for k, v in staged.items()}


_FIELD_WORDS = {"amount": "amount", "currency": "currency", "status": "status"}


# --- Preview ---

async def preview_budget_sheet(filename: str, data: bytes, actor_id: UUID, year: int, currency: str) -> dict:
    """Check every row of an uploaded budget tracker and keep the rows that pass for confirming later."""
    try:
        currency = parse_currency(currency, DEFAULT_CURRENCY)
    except InvalidValueError as exc:
        raise ImportFileError(str(exc)) from exc
    known = [name for names in COLUMN_ALIASES.values() for name in names]
    rows = read_table(filename, data, known, sheet_names=SHEET_NAMES)
    if not rows:
        raise ImportFileError("The sheet has a header but no budget rows.")
    headers = [h for h in rows[0] if h != ROW_NUMBER]
    match = match_columns(headers, COLUMN_ALIASES)
    labels = (("month", "Month"), ("brand", "Brand"), ("market", "Market / Region"), ("amount", "Approved Budget"))
    missing = [label for field, label in labels if field not in match.fields]
    if missing:
        plural = "s" if len(missing) > 1 else ""
        raise ImportFileError(f"The sheet needs {_join_words(missing)} column{plural}. Start from the template.")
    amounts = [row.get(match.fields["amount"], "") for row in rows]

    pool = get_pool()
    async with pool.acquire() as conn:
        brands = await conn.fetch("SELECT id, code, name FROM brands WHERE is_active")
        markets = await conn.fetch("SELECT id, code, name FROM markets WHERE is_active")
        channels = await conn.fetch("SELECT id, code, name FROM channels WHERE is_active")
        ctx = _Context(match.fields, uses_decimal_comma(amounts), year, currency, brands, markets, channels)
        previews = [_check_row(row, ctx) for row in rows]

        months = sorted({p["_staged"]["month"] for p in previews if p["_staged"]})
        existing_rows = await conn.fetch(
            "SELECT id, month, brand_id, market_id, channel_id, amount, currency, status "
            "FROM monthly_budgets WHERE month = ANY($1::date[])",
            months,
        )
        existing = {_key(r): dict(r) for r in existing_rows}
        notes = await conn.fetch(
            "SELECT budget_id, body FROM budget_comments WHERE budget_id = ANY($1::uuid[])",
            [r["id"] for r in existing_rows],
        )
        known_notes = {(n["budget_id"], n["body"]) for n in notes}

        seen: dict[tuple, int] = {}
        staged_rows = []
        for p in previews:
            staged = p.pop("_staged")
            if staged is None:
                continue
            key = _key(staged)
            if key in seen:
                p["status"] = "INVALID"
                p["messages"] = [f"Same month, brand, market and channel as row {seen[key]}.", *p["messages"]]
                continue
            seen[key] = p["row_number"]
            current = existing.get(key)
            if current is not None:
                p["budget_id"] = current["id"]
                changed = _changed(current, staged)
                new_note = staged["note"] and (current["id"], staged["note"]) not in known_notes
                if changed:
                    p["status"] = "UPDATE"
                    p["messages"] = [f"Updates the {_join_words([_FIELD_WORDS[c] for c in changed])}.", *p["messages"]]
                elif new_note:
                    p["status"] = "UPDATE"
                else:
                    p["status"] = "UNCHANGED"
                if new_note:
                    p["messages"] = [*p["messages"], "Adds the note to the budget."]
            # Unchanged rows are kept too: the budget may have changed by the time the upload is confirmed
            staged_rows.append(_json_ready(staged))

        counts = {
            "new": sum(p["status"] == "NEW" for p in previews),
            "update": sum(p["status"] == "UPDATE" for p in previews),
            "unchanged": sum(p["status"] == "UNCHANGED" for p in previews),
            "invalid": sum(p["status"] == "INVALID" for p in previews),
        }
        batch_id = await conn.fetchval(
            """
            INSERT INTO import_batches (kind, filename, status, total_rows, ready_rows, error_rows, duplicate_rows,
                                        staged_rows, created_by)
            VALUES ($1, $2, 'PREVIEW', $3, $4, $5, 0, $6::jsonb, $7)
            RETURNING id
            """,
            KIND, (filename or "upload")[:255], len(previews), counts["new"] + counts["update"], counts["invalid"],
            json.dumps(staged_rows), actor_id,
        )
    return {
        "batch_id": batch_id, "filename": filename, "rows": previews, "counts": counts,
        "ignored_columns": [" ".join(h.split()) for h in match.ignored],
    }


# --- Confirm and undo ---

def _is_expired(created_at: datetime) -> bool:
    return created_at < datetime.now(timezone.utc) - PREVIEW_LIFETIME


def _from_json(staged: dict) -> dict:
    return {**staged, "month": date.fromisoformat(staged["month"]), "amount": Decimal(staged["amount"])}


def _audit_values(row: dict) -> dict:
    return {
        "month": row["month"], "brand_id": row["brand_id"], "market_id": row["market_id"],
        "channel_id": row["channel_id"], "amount": float(row["amount"]), "currency": row["currency"],
        "status": row["status"],
    }


async def confirm_import(batch_id: UUID, actor_id: UUID) -> dict:
    """Add and update the previewed budgets and their notes, then mark the upload confirmed."""
    expired = False
    pool = get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            batch = await conn.fetchrow(
                "SELECT filename, status, staged_rows, created_by, created_at FROM import_batches "
                "WHERE id = $1 AND kind = $2 FOR UPDATE",
                batch_id, KIND,
            )
            if batch is None:
                raise ImportNotFoundError
            if batch["created_by"] != actor_id:
                raise ImportForbiddenError
            if batch["status"] != "PREVIEW":
                raise ImportConflictError(f"This upload is already {batch['status'].lower()}.")
            if _is_expired(batch["created_at"]):
                await conn.execute("UPDATE import_batches SET status = 'EXPIRED' WHERE id = $1", batch_id)
                expired = True
            else:
                raw = json.loads(batch["staged_rows"]) if isinstance(batch["staged_rows"], str) else batch["staged_rows"]
                summary = await _apply(conn, [_from_json(r) for r in raw], actor_id, batch_id)
                await log_audit(
                    conn, actor_id, "IMPORT", AUDIT_ENTITY, batch_id,
                    new_value={"kind": KIND, "filename": batch["filename"],
                               **{k: v for k, v in summary.items() if k != "undo_data"}},
                )
                undo_data = summary.pop("undo_data")
                await conn.execute(
                    """
                    UPDATE import_batches
                    SET status = 'CONFIRMED', imported_rows = $2, confirmed_at = now(), confirmed_by = $3,
                        undo_data = $4::jsonb
                    WHERE id = $1
                    """,
                    batch_id, summary["created"] + summary["updated"], actor_id, json.dumps(undo_data, default=str),
                )
    if expired:
        raise ImportExpiredError
    return {"batch_id": batch_id, "status": "CONFIRMED", **summary}


async def _apply(conn, staged: list[dict], actor_id: UUID, batch_id: UUID) -> dict:
    created: list[dict] = []
    updated: list[dict] = []
    note_ids: list[str] = []
    unchanged = skipped = 0
    for row in sorted(staged, key=lambda r: r["row_number"]):
        current = await conn.fetchrow(
            """
            SELECT id, amount, currency, status FROM monthly_budgets
            WHERE month = $1 AND COALESCE(brand_id, 0) = $2 AND market_id = $3 AND COALESCE(channel_id, 0) = $4
            FOR UPDATE
            """,
            row["month"], row["brand_id"] or 0, row["market_id"], row["channel_id"] or 0,
        )
        values = _audit_values(row)
        if current is None:
            budget_id = await conn.fetchval(
                """
                INSERT INTO monthly_budgets (month, brand_id, market_id, channel_id, amount, currency, status,
                                             created_by, updated_by)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $8)
                RETURNING id
                """,
                row["month"], row["brand_id"], row["market_id"], row["channel_id"], row["amount"], row["currency"],
                row["status"], actor_id,
            )
            await log_audit(conn, actor_id, "CREATE", BUDGET_ENTITY, budget_id, new_value={**values, "import": str(batch_id)})
            created.append({"id": str(budget_id), "set": {c: str(row[c]) for c in _VALUE_COLUMNS}})
        else:
            budget_id = current["id"]
            changed = _changed(dict(current), row)
            if "currency" in changed and await conn.fetchval(
                "SELECT EXISTS (SELECT 1 FROM budget_transfers WHERE from_budget_id = $1 OR to_budget_id = $1)", budget_id,
            ):
                skipped += 1  # changing the currency would break its money moves
                continue
            if changed:
                await conn.execute(
                    "UPDATE monthly_budgets SET amount = $2, currency = $3, status = $4, updated_at = now(), "
                    "updated_by = $5 WHERE id = $1",
                    budget_id, row["amount"], row["currency"], row["status"], actor_id,
                )
                previous = {c: str(current[c]) for c in _VALUE_COLUMNS}
                await log_audit(
                    conn, actor_id, "UPDATE", BUDGET_ENTITY, budget_id,
                    old_value={**values, **{c: current[c] for c in changed}, "amount": float(current["amount"])},
                    new_value={**values, "import": str(batch_id)},
                )
                updated.append({"id": str(budget_id), "previous": previous, "set": {c: str(row[c]) for c in _VALUE_COLUMNS}})
            else:
                unchanged += 1
        if row["note"]:
            already = await conn.fetchval(
                "SELECT EXISTS (SELECT 1 FROM budget_comments WHERE budget_id = $1 AND body = $2)", budget_id, row["note"],
            )
            if not already:
                comment_id = await conn.fetchval(
                    "INSERT INTO budget_comments (budget_id, month, body, author_id) VALUES ($1, $2, $3, $4) RETURNING id",
                    budget_id, row["month"], row["note"], actor_id,
                )
                await log_audit(
                    conn, actor_id, "CREATE", COMMENT_ENTITY, comment_id,
                    new_value={"budget_id": budget_id, "month": row["month"], "body": row["note"], "import": str(batch_id)},
                )
                note_ids.append(str(comment_id))
    return {
        "created": len(created), "updated": len(updated), "unchanged": unchanged, "notes_added": len(note_ids),
        "skipped": skipped, "undo_data": {"created": created, "updated": updated, "notes": note_ids},
    }


def _same(current, values: dict) -> bool:
    """True while a budget still has the values the upload gave it."""
    return (
        Decimal(current["amount"]) == Decimal(values["amount"])
        and current["currency"] == values["currency"]
        and current["status"] == values["status"]
    )


async def undo_import(batch_id: UUID, actor_id: UUID) -> dict:
    """Remove the budgets and notes an upload added and restore the previous values of the budgets it changed.

    Budgets changed again since the upload, budgets it added that money moves now use, and budgets whose
    currency it changed that money moves now use, are left as they are.
    """
    pool = get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            batch = await conn.fetchrow(
                "SELECT filename, status, undo_data FROM import_batches WHERE id = $1 AND kind = $2 FOR UPDATE",
                batch_id, KIND,
            )
            if batch is None:
                raise ImportNotFoundError
            if batch["status"] != "CONFIRMED":
                raise ImportConflictError("Only a confirmed upload can be undone.")
            data = batch["undo_data"]
            data = json.loads(data) if isinstance(data, str) else (data or {})

            notes = await conn.fetch(
                "DELETE FROM budget_comments WHERE id = ANY($1::uuid[]) RETURNING id, budget_id, month, body",
                [UUID(i) for i in data.get("notes", [])],
            )
            for n in notes:
                await log_audit(conn, actor_id, "DELETE", COMMENT_ENTITY, n["id"],
                                old_value={"budget_id": n["budget_id"], "month": n["month"], "body": n["body"]})

            removed = restored = skipped = 0
            for item in data.get("created", []):
                current = await conn.fetchrow(
                    "SELECT id, month, brand_id, market_id, channel_id, amount, currency, status FROM monthly_budgets "
                    "WHERE id = $1 FOR UPDATE",
                    UUID(item["id"]),
                )
                if current is None:
                    continue  # already deleted
                in_use = await conn.fetchval(
                    "SELECT EXISTS (SELECT 1 FROM budget_transfers WHERE from_budget_id = $1 OR to_budget_id = $1)", current["id"],
                )
                if in_use or not _same(current, item["set"]):
                    skipped += 1
                    continue
                await conn.execute("DELETE FROM monthly_budgets WHERE id = $1", current["id"])
                await log_audit(conn, actor_id, "DELETE", BUDGET_ENTITY, current["id"], old_value=_audit_values(dict(current)))
                removed += 1
            for item in data.get("updated", []):
                current = await conn.fetchrow(
                    "SELECT id, month, brand_id, market_id, channel_id, amount, currency, status FROM monthly_budgets "
                    "WHERE id = $1 FOR UPDATE",
                    UUID(item["id"]),
                )
                if current is None:
                    continue  # deleted since
                if not _same(current, item["set"]):
                    skipped += 1
                    continue
                previous = item["previous"]
                if previous["currency"] != current["currency"] and await conn.fetchval(
                    "SELECT EXISTS (SELECT 1 FROM budget_transfers WHERE from_budget_id = $1 OR to_budget_id = $1)",
                    current["id"],
                ):
                    skipped += 1  # changing the currency back would break the money moves made since
                    continue
                await conn.execute(
                    "UPDATE monthly_budgets SET amount = $2, currency = $3, status = $4, updated_at = now(), "
                    "updated_by = $5 WHERE id = $1",
                    current["id"], Decimal(previous["amount"]), previous["currency"], previous["status"], actor_id,
                )
                await log_audit(
                    conn, actor_id, "UPDATE", BUDGET_ENTITY, current["id"],
                    old_value=_audit_values(dict(current)),
                    new_value={**_audit_values(dict(current)), **previous, "amount": float(previous["amount"])},
                )
                restored += 1

            await conn.execute(
                "UPDATE import_batches SET status = 'UNDONE', undone_at = now(), undone_by = $2 WHERE id = $1",
                batch_id, actor_id,
            )
            result = {"removed": removed, "restored": restored, "notes_removed": len(notes), "skipped": skipped}
            await log_audit(
                conn, actor_id, "UNDO_IMPORT", AUDIT_ENTITY, batch_id,
                old_value={"kind": KIND, "filename": batch["filename"], **result},
            )
    return {"batch_id": batch_id, "status": "UNDONE", **result}


# --- Template ---

async def build_template() -> bytes:
    """An .xlsx with the Budget Tracker header, and dropdowns for the brand, market and spend progress."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    pool = get_pool()
    async with pool.acquire() as conn:
        brands = await conn.fetch("SELECT name FROM brands WHERE is_active ORDER BY sort_order, name")
        markets = await conn.fetch("SELECT name FROM markets WHERE is_active ORDER BY sort_order, name")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAMES[0]
    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EAF6")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    widths = {"Notes": 50, "Market / Region": 18, "Approved Budget": 16, "Spend Progress": 16}
    for index, header in enumerate(TEMPLATE_COLUMNS, start=1):
        column = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[column].width = widths.get(header, 18)

    # Lists live on a hidden sheet so long brand lists still fit Excel's dropdown limits
    lists = workbook.create_sheet("Lists")
    lists.sheet_state = "hidden"
    progress = ["Complete", "In Progress", "Not Started", "Not committed", "Canceled"]
    for column, values in enumerate(([b["name"] for b in brands], [m["name"] for m in markets], progress), start=1):
        for row, value in enumerate(values, start=1):
            lists.cell(row=row, column=column, value=value)
    for letter, target, count in (("A", "B", len(brands)), ("B", "C", len(markets)), ("C", "E", len(progress))):
        validation = DataValidation(type="list", formula1=f"=Lists!${letter}$1:${letter}${max(count, 1)}", allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{target}2:{target}1000")
    for row in range(2, 1001):
        sheet.cell(row=row, column=4).number_format = '"$"#,##0'

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
