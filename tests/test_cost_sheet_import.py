"""Cost sheet uploads: preview, confirm, undo, template and history. Sheets are built in memory."""
import io
import json
import uuid
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from services.cost_sheet_import_service import TEMPLATE_COLUMNS
from tests.conftest import _delete_user, _make_user, _token_for, run_sql

HEADER = ["Item", "Category", "Vendor", "Estimate", "Actual", "Status", "Invoice #", "Date", "Paid On"]


def _xlsx(rows: list[list], title: str | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if title:
        sheet.append([title])
        sheet.append([])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def event(client, admin, ref):
    response = client.post(
        "/api/initiatives",
        json={
            "name": f"Pytest Upload Expo {uuid.uuid4().hex[:6]}",
            "brand_id": ref["brands"]["DESKFLEX"],
            "channel_id": ref["channels"]["EVENTS"],
            "sub_channel_id": ref["sub_channels"]["IN_PERSON_EVENTS"],
            "currency": "USD",
            "start_date": "2026-11-04",
            "end_date": "2026-11-06",
            "event": {"city": "Las Vegas", "country_code": "US"},
        },
        headers=admin["headers"],
    )
    assert response.status_code == 201, response.text
    return response.json()


@pytest.fixture
def other_admin(client):
    """A second temporary super admin."""
    user_id = _make_user(super_admin=True)
    try:
        yield {"id": user_id, "headers": _token_for(client, user_id)}
    finally:
        _delete_user(user_id)


@pytest.fixture(scope="module")
def categories():
    return {r["code"]: r["id"] for r in run_sql("SELECT id, code FROM spend_categories")}


def _upload(client, user, initiative_id, data: bytes, filename="costs.xlsx"):
    return client.post(
        f"/api/initiatives/{initiative_id}/imports/cost-sheet",
        files={"file": (filename, data, "application/octet-stream")},
        headers=user["headers"],
    )


def _confirm(client, user, batch_id):
    return client.post(f"/api/imports/{batch_id}/confirm", headers=user["headers"])


def _rows(preview) -> dict[int, dict]:
    return {row["row_number"]: row for row in preview["rows"]}


# --- Preview ---

def test_estimates_only_with_a_title_row(client, admin, event):
    sheet = _xlsx(
        [
            HEADER,
            ["10x20 booth space", "Booth space", "Show organiser", 8000, None, None, None, datetime(2026, 6, 15), None],
            ["Carpet", "carpet & flooring", "Expo Services", "$600", None, None, None, "04/11/2026", None],
        ],
        title="Facility Expo — cost sheet",
    )

    response = _upload(client, admin, event["id"], sheet)

    assert response.status_code == 201, response.text
    preview = response.json()
    assert preview["counts"] == {"total": 2, "ready": 2, "warnings": 0, "duplicates": 0, "invalid": 0}
    rows = _rows(preview)
    assert set(rows) == {4, 5}  # sheet row numbers, below the title and blank row
    assert (rows[4]["status"], rows[4]["status_value"], rows[4]["category"], rows[4]["expense_type"]) == (
        "READY", "Estimate", "Booth space", "Space & sponsorship",
    )
    assert (rows[5]["category"], rows[5]["estimate"], rows[5]["date"]) == ("Carpet & flooring", 600, "2026-11-04")
    assert preview["totals_preview"] == {"expected_total": 8600, "committed": 0, "actual": 0}


def test_statuses_are_inferred_from_actuals_and_paid_dates(client, admin, event):
    sheet = _xlsx([
        HEADER,
        ["Booth space", "Booth space", "Organiser", 8000, 8000, None, "INV-A1", "2026-06-15", "2026-07-01"],
        ["Booth build", "Booth build or rental", "ExhibitWorks", 4000, 4500, None, "INV-A2", "2026-10-01", None],
        ["Lighting", "Lighting", "Expo Services", 450, None, "committed", None, "2026-11-04", None],
    ])

    preview = _upload(client, admin, event["id"], sheet).json()

    rows = _rows(preview)
    assert [rows[n]["status_value"] for n in (2, 3, 4)] == ["Paid", "Invoiced", "Committed"]
    assert all(row["status"] == "READY" for row in preview["rows"])
    assert preview["totals_preview"] == {"expected_total": 12950, "committed": 12950, "actual": 12500}


def test_aliases_and_ignored_columns(client, admin, event):
    sheet = _xlsx([
        ["Description", "Category", "Supplier", "Quote", "Internal", "Invoice Number"],
        ["Freight", "Freight", "Expo Freight", 1800, "x", "F-1"],
    ])

    preview = _upload(client, admin, event["id"], sheet).json()

    row = preview["rows"][0]
    assert (row["item"], row["vendor"], row["estimate"], row["invoice_ref"]) == ("Freight", "Expo Freight", 1800, "F-1")
    assert preview["ignored_columns"] == ["Internal"]


def test_receipt_columns(client, admin, event):
    """A card purchase at the show comes with a receipt, not an invoice; its number and date are read the same way."""
    sheet = _xlsx([
        ["Item", "Category", "Vendor", "Actual", "Receipt #", "Receipt date", "Paid On", "Payment Method"],
        ["Exhibitor badges", "Badges", "Expo Organisers", 300, "R-5521", "2026-11-04", "2026-11-04", "Corporate card"],
    ])

    preview = _upload(client, admin, event["id"], sheet).json()

    row = preview["rows"][0]
    assert (row["invoice_ref"], row["date"], row["status_value"]) == ("R-5521", "2026-11-04", "Paid")
    assert preview["ignored_columns"] == []


def test_invalid_rows_list_every_problem(client, admin, event):
    sheet = _xlsx([
        ["Item", "Category", "Expense type", "Estimate", "Status", "Date"],
        ["Mystery", "Booth stuff", None, 100, None, "2026-11-04"],
        ["Flights to Vegas", "Flights", None, 1200, None, "2026-11-01"],
        ["Lighting", "Lighting", "Shipping & logistics", "#N/A", "Sorted", "not a date"],
        ["Nothing", "Lighting", None, None, None, "2026-11-04"],
        ["Ads", "Media / ads", None, 50, None, "2026-11-04"],
    ])

    preview = _upload(client, admin, event["id"], sheet).json()

    rows = _rows(preview)
    assert all(row["status"] == "INVALID" for row in preview["rows"])
    assert rows[2]["messages"] == ["Unknown category 'Booth stuff' — choose one from the template's Categories sheet"]
    assert rows[3]["messages"] == ["Travel comes from the travel management system — leave it out of the sheet."]
    assert rows[4]["messages"] == [
        "Lighting is under Show services, not Shipping & logistics",
        "Estimate: #N/A is a spreadsheet error, not an amount",
        "Date: 'not a date' isn't a date. Use DD/MM/YYYY or YYYY-MM-DD.",
        "Unknown status 'Sorted' — use Estimate, Committed, Invoiced or Paid",
    ]
    assert rows[5]["messages"] == ["Add an estimate or an actual amount"]
    assert rows[6]["messages"] == ["Media / ads: That category isn't used for this channel."]
    assert preview["counts"]["invalid"] == 5 and preview["counts"]["ready"] == 0


def test_paid_without_a_paid_date_is_invalid(client, admin, event):
    sheet = _xlsx([HEADER, ["Badges", "Badges & passes", "Organiser", 300, 300, "Paid", None, "2026-09-15", None]])

    row = _upload(client, admin, event["id"], sheet).json()["rows"][0]

    assert (row["status"], row["messages"]) == ("INVALID", ["Paid needs a Paid On date."])


def test_duplicates_against_existing_costs_and_within_the_file(client, admin, event, categories):
    existing = client.post(
        f"/api/initiatives/{event['id']}/spend-lines",
        json={"category_id": categories["FLOORING"], "amount": 600, "incurred_on": "2026-11-04", "vendor": "Freeman", "invoice_ref": "INV-77"},
        headers=admin["headers"],
    )
    assert existing.status_code == 201, existing.text
    sheet = _xlsx([
        HEADER,
        ["Carpet again", "Carpet & flooring", "FREEMAN", None, 600, None, "INV-77", "2026-11-04", None],
        ["Lighting", "Lighting", "GES", None, 450, None, "G-1", "2026-11-04", None],
        ["Lighting (twice)", "Lighting", "ges", None, 450, None, "G-1", "2026-11-04", None],
    ])

    preview = _upload(client, admin, event["id"], sheet).json()

    rows = _rows(preview)
    assert rows[2]["status"] == "DUPLICATE"
    assert rows[2]["messages"] == [f"Invoice INV-77 from Freeman is already recorded on {event['event_code']} · {event['name']}."]
    assert rows[3]["status"] == "READY"
    assert (rows[4]["status"], rows[4]["messages"]) == ("DUPLICATE", ["Invoice G-1 from ges appears twice in this file (row 3)."])
    assert preview["counts"]["duplicates"] == 2


def test_missing_date_uses_the_event_start_with_a_warning(client, admin, event):
    sheet = _xlsx([["Item", "Category", "Estimate"], ["Freight", "Freight", 1800]])

    row = _upload(client, admin, event["id"], sheet).json()["rows"][0]

    assert (row["status"], row["date"], row["messages"]) == (
        "READY_WITH_WARNINGS", "2026-11-04", ["No date — used the event's start date"],
    )


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("costs.pdf", b"%PDF", "Upload an Excel (.xlsx) or CSV file."),
        ("costs.xlsx", None, "The sheet needs Category column. Start from the template."),
    ],
)
def test_unusable_files(client, admin, event, filename, data, message):
    data = data or _xlsx([["Item", "Estimate", "Vendor"], ["Freight", 100, "X"]])

    response = _upload(client, admin, event["id"], data, filename)

    assert (response.status_code, response.json()["detail"]) == (422, message)


def test_viewer_cannot_upload(client, viewer, event):
    response = _upload(client, viewer, event["id"], _xlsx([HEADER]))

    assert response.status_code == 403


# --- Confirm ---

def _ready_sheet() -> bytes:
    return _xlsx([
        HEADER,
        ["Booth space", "Booth space", "Organiser", 8000, 8000, None, "UP-1", "2026-06-15", "2026-07-01"],
        ["Freight", "Freight", "Expo Freight", 1800, None, None, None, "2026-10-28", None],
        ["Mystery", "Booth stuff", None, 100, None, None, None, "2026-11-04", None],  # invalid, not imported
    ])


def test_confirm_creates_cost_lines(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()

    response = _confirm(client, admin, preview["batch_id"])

    assert response.status_code == 200, response.text
    assert response.json() == {"batch_id": preview["batch_id"], "status": "CONFIRMED", "imported_rows": 2, "skipped": 0}
    lines = run_sql(
        "SELECT description, status, amount, estimated_amount, source, import_batch_id, created_by FROM spend_lines "
        "WHERE initiative_id = $1 ORDER BY description",
        event["id"],
    )
    assert [(r["description"], r["status"], float(r["amount"])) for r in lines] == [("Booth space", "PAID", 8000), ("Freight", "ESTIMATED", 1800)]
    assert all(r["source"] == "COST_SHEET_UPLOAD" and str(r["import_batch_id"]) == preview["batch_id"] for r in lines)
    assert float(lines[1]["estimated_amount"]) == 1800
    audit = run_sql("SELECT action, new_value FROM audit_logs WHERE entity_type = 'ImportBatch' AND entity_id = $1", preview["batch_id"])
    assert audit[0]["action"] == "IMPORT"
    assert json.loads(audit[0]["new_value"])["imported_rows"] == 2
    # The lines are ordinary costs: the cost sheet and its summary include them
    summary = client.get(f"/api/initiatives/{event['id']}/spend-summary", headers=admin["headers"]).json()
    assert summary["totals"]["expected_total"] == 9800


def test_confirm_skips_invoices_recorded_since_the_preview(client, admin, event, categories):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()
    client.post(
        f"/api/initiatives/{event['id']}/spend-lines",
        json={"category_id": categories["BOOTH_SPACE"], "amount": 8000, "incurred_on": "2026-06-15", "vendor": "organiser", "invoice_ref": "UP-1"},
        headers=admin["headers"],
    )

    response = _confirm(client, admin, preview["batch_id"])

    assert (response.json()["imported_rows"], response.json()["skipped"]) == (1, 1)


def test_confirming_twice_conflicts(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()
    _confirm(client, admin, preview["batch_id"])

    response = _confirm(client, admin, preview["batch_id"])

    assert (response.status_code, response.json()["detail"]) == (409, "This upload is already confirmed.")


def test_only_the_uploader_can_confirm(client, admin, other_admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()

    response = _confirm(client, other_admin, preview["batch_id"])

    assert response.status_code == 403
    assert run_sql("SELECT status FROM import_batches WHERE id = $1", preview["batch_id"])[0]["status"] == "PREVIEW"


def test_expired_preview(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()
    run_sql("UPDATE import_batches SET created_at = now() - interval '25 hours' WHERE id = $1", preview["batch_id"])

    response = _confirm(client, admin, preview["batch_id"])

    assert (response.status_code, response.json()["detail"]) == (410, "This preview has expired. Upload the file again.")
    assert run_sql("SELECT status FROM import_batches WHERE id = $1", preview["batch_id"])[0]["status"] == "EXPIRED"
    assert run_sql("SELECT 1 FROM spend_lines WHERE initiative_id = $1", event["id"]) == []


def test_confirm_missing_batch(client, admin):
    assert _confirm(client, admin, uuid.uuid4()).status_code == 404


# --- Undo ---

def test_undo_removes_the_lines(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()
    _confirm(client, admin, preview["batch_id"])

    response = client.post(f"/api/imports/{preview['batch_id']}/undo", headers=admin["headers"])

    assert response.status_code == 200, response.text
    assert response.json() == {"batch_id": preview["batch_id"], "status": "UNDONE", "removed_rows": 2}
    assert run_sql("SELECT 1 FROM spend_lines WHERE initiative_id = $1", event["id"]) == []
    assert run_sql("SELECT action FROM audit_logs WHERE entity_id = $1 AND action = 'UNDO_IMPORT'", preview["batch_id"])
    # Undo only works once
    again = client.post(f"/api/imports/{preview['batch_id']}/undo", headers=admin["headers"])
    assert again.status_code == 409


def test_undo_blocked_by_a_reconciled_line(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()
    _confirm(client, admin, preview["batch_id"])
    run_sql(
        "UPDATE spend_lines SET reconciled_month = $2 WHERE import_batch_id = $1 AND status = 'PAID'",
        uuid.UUID(preview["batch_id"]), date(2026, 7, 1),
    )

    response = client.post(f"/api/imports/{preview['batch_id']}/undo", headers=admin["headers"])

    assert (response.status_code, response.json()["detail"]) == (
        409, "1 of these costs is already reconciled — ask Finance to reopen them first.",
    )
    assert len(run_sql("SELECT 1 FROM spend_lines WHERE initiative_id = $1", event["id"])) == 2


def test_undo_needs_a_confirmed_upload(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()

    response = client.post(f"/api/imports/{preview['batch_id']}/undo", headers=admin["headers"])

    assert response.status_code == 409


# --- History and template ---

def test_history_lists_uploads(client, admin, event):
    first = _upload(client, admin, event["id"], _ready_sheet(), "first.xlsx").json()
    _confirm(client, admin, first["batch_id"])
    _upload(client, admin, event["id"], _ready_sheet(), "second.csv.xlsx")

    history = client.get(f"/api/initiatives/{event['id']}/imports", headers=admin["headers"]).json()

    assert [h["filename"] for h in history] == ["second.csv.xlsx", "first.xlsx"]
    confirmed = history[1]
    assert (confirmed["status"], confirmed["imported_rows"], confirmed["line_count"], confirmed["error_rows"]) == ("CONFIRMED", 2, 2, 1)
    assert confirmed["created_by_name"] == "Pytest User"


def test_template(client, admin):
    response = client.get("/api/imports/cost-sheet/template", headers=admin["headers"])

    assert response.status_code == 200
    assert 'filename="event-cost-sheet-template.xlsx"' in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    costs, categories = workbook["Costs"], workbook["Categories"]
    assert [c.value for c in costs[1]] == TEMPLATE_COLUMNS
    assert [costs.cell(row=r, column=2).value for r in range(2, 10)] == [
        "Booth space", "Booth build or rental", "Carpet & flooring", "Lighting", "Electrical", "Badges & passes", "Freight", "Printing & signage",
    ]
    category_names = {row[1] for row in categories.iter_rows(min_row=2, values_only=True)}
    assert "Lighting" in category_names
    assert not category_names & {"Flights", "Hotels", "Ground transport", "Per diems & meals", "Other travel", "Media / ads"}
    dropdowns = {str(v.sqref): v.formula1 for v in costs.data_validations.dataValidation}
    assert dropdowns["B2:B1000"].startswith("Categories!$B$2:$B$") or dropdowns["B2:B1000"].startswith("=Categories")
    assert dropdowns["G2:G1000"] == '"Estimate,Committed,Invoiced,Paid"'
    assert dropdowns["K2:K1000"] == '"AP invoice,Corporate card,Other"'
    assert dropdowns["M2:M1000"] == '"OPEX,CAPEX"'


def test_template_round_trips(client, admin, event):
    """The template's own example rows upload cleanly."""
    template = client.get("/api/imports/cost-sheet/template", headers=admin["headers"]).content

    preview = _upload(client, admin, event["id"], template, "event-cost-sheet-template.xlsx").json()

    assert preview["counts"]["total"] == 8
    assert preview["counts"]["invalid"] == 0, [r for r in preview["rows"] if r["status"] == "INVALID"]


def test_amounts_over_the_limit_are_invalid(client, admin, event):
    sheet = _xlsx([HEADER, ["Huge", "Freight", "Expo Freight", 1e15, None, None, None, "2026-10-28", None]])

    row = _upload(client, admin, event["id"], sheet).json()["rows"][0]

    assert (row["status"], row["messages"]) == ("INVALID", ["Estimate: too large"])


def test_each_imported_and_undone_line_is_audited(client, admin, event):
    preview = _upload(client, admin, event["id"], _ready_sheet()).json()
    _confirm(client, admin, preview["batch_id"])
    line_ids = [r["id"] for r in run_sql("SELECT id FROM spend_lines WHERE import_batch_id = $1", preview["batch_id"])]

    client.post(f"/api/imports/{preview['batch_id']}/undo", headers=admin["headers"])

    audit = run_sql(
        "SELECT action, old_value FROM audit_logs WHERE entity_type = 'SpendLine' AND entity_id = ANY($1::text[]) ORDER BY created_at",
        [str(i) for i in line_ids],
    )
    assert sorted(a["action"] for a in audit) == ["CREATE", "CREATE", "DELETE", "DELETE"]
    assert all(json.loads(a["old_value"])["description"] for a in audit if a["action"] == "DELETE")
