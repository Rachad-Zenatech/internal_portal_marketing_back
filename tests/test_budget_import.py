"""Budget tracker upload: /api/budget-imports (preview, confirm, undo) and the template.

Budgets here are in 2093, which no real data uses; they're removed after each test.
"""
import io
import uuid
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from services.budget_import_service import TEMPLATE_COLUMNS, parse_month
from services.import_parsing import InvalidValueError
from tests.conftest import _delete_user, _make_user, _token_for, run_sql

YEAR = 2093
_TEST_MONTHS = f"month >= DATE '{YEAR}-01-01' AND month < DATE '{YEAR + 1}-01-01'"

# Marketing's "Budget Tracker" header exactly as in their sheet (line breaks included)
SHEET_HEADER = [
    "Month", "Brand", "Market / Region", "Approved Budget", "Spend Progress", "Actual Spend to Date",
    "Carryover In\n(from prior month)", "Unspent\n(carries to next month)", "Carryover Spent\nThis Month", "Notes",
    "Last Update",
]


@pytest.fixture(autouse=True)
def _clean_budgets():
    yield
    run_sql(
        f"DELETE FROM budget_transfers WHERE from_budget_id IN (SELECT id FROM monthly_budgets WHERE {_TEST_MONTHS}) "
        f"OR to_budget_id IN (SELECT id FROM monthly_budgets WHERE {_TEST_MONTHS})"
    )
    run_sql(f"DELETE FROM budget_comments WHERE {_TEST_MONTHS}")
    run_sql(f"DELETE FROM monthly_budgets WHERE {_TEST_MONTHS}")


@pytest.fixture
def other_admin(client):
    user_id = _make_user(super_admin=True)
    yield {"id": user_id, "headers": _token_for(client, user_id)}
    _delete_user(user_id)


@pytest.fixture(scope="module")
def markets():
    return {r["code"]: r["id"] for r in run_sql("SELECT id, code FROM markets")}


def _xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    # A different tab first: the upload finds the Budget Tracker tab by name
    workbook.active.title = "Summary"
    workbook.active.append(["Total", 100])
    sheet = workbook.create_sheet("Budget Tracker")
    sheet.append(["Marketing budget 2093"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, user, data: bytes, filename="budgets.xlsx", year=YEAR, currency="USD"):
    return client.post(
        "/api/budget-imports",
        files={"file": (filename, data, "application/octet-stream")},
        data={"year": str(year), "currency": currency},
        headers=user["headers"],
    )


def _sheet() -> bytes:
    """Marketing's August and September rows, as pasted from their tracker."""
    return _xlsx([
        SHEET_HEADER,
        # 3
        ["Aug", "Othership", "London", 10000, "Complete", 13000, None, None, 3000, "Spent on SEO team ads", None],
        # 4
        ["Aug", "TBD/All", "Unassigned", 10000, "In Progress", 3000, None, 7000, None, "Spent $3k on Othership", None],
        # 5: paused, with a last update date
        ["Sep", "Othership", "London", 10000, "Complete", 11650, None, 5350, 1650, "Campaign paused Sep 21", datetime(YEAR, 9, 22)],
        # 6: typed as text with a symbol
        ["Sep", "WorkAware", "North America", "$25,000", "In Progress", 7800, None, " ", " ", " = 11k cad g ads- ", None],
        # 7: no approved budget
        ["Sep", "DeskFlex", "North America", None, "In Progress", 12000, None, None, None, "9k g-ads. 3k bing", None],
        # 8: cancelled, placeholders everywhere
        ["Sep", "Othership", "North America", None, "Canceled", ".", ".", ".", ".", "Canceled", None],
        # 9: pounds
        ["Sep", "Tillerstack", "Germany", "£10,000", "Not Started", None, None, 10000, None, "Anja to follow up", None],
        # 10
        ["Sep", "TBD/All", "London", 10000, "Not committed", None, None, None, None, None, None],
        # 11: unknown brand, market and progress
        ["Sep", "Acme", "Mars", 10, "Sleeping", None, None, None, None, None, None],
        # 12: not a month
        ["Smarch", "DeskFlex", "Korea", 10, "Complete", None, None, None, None, None, None],
        # 13: row 3 again
        ["August", "Othership", "London", 1, "Complete", None, None, None, None, None, None],
    ])


def test_parse_month():
    assert parse_month("Aug", 2093) == date(2093, 8, 1)
    assert parse_month("September", 2093) == date(2093, 9, 1)
    assert parse_month("Sept", 2093) == date(2093, 9, 1)
    assert parse_month("Sep 2094", 2093) == date(2094, 9, 1)
    assert parse_month("Sep-94", 2093) == date(2094, 9, 1)
    assert parse_month("2094-03", 2093) == date(2094, 3, 1)
    assert parse_month("03/2094", 2093) == date(2094, 3, 1)
    assert parse_month("2094-03-15", 2093) == date(2094, 3, 1)
    with pytest.raises(InvalidValueError):
        parse_month("Smarch", 2093)


def test_preview(client, admin):
    response = _upload(client, admin, _sheet())

    assert response.status_code == 201, response.text
    preview = response.json()
    rows = {r["row_number"]: r for r in preview["rows"]}
    assert preview["counts"] == {"new": 8, "update": 0, "unchanged": 0, "invalid": 3}
    assert "Actual Spend to Date" in preview["ignored_columns"]
    assert "Carryover In (from prior month)" in preview["ignored_columns"]

    assert rows[3] == {
        "row_number": 3, "status": "NEW", "messages": [], "month": f"{YEAR}-08-01", "brand": "Othership", "market": "London",
        "channel": None, "amount": 10000.0, "currency": "USD", "budget_status": "APPROVED", "sheet_progress": "Complete",
        "note": "Spent on SEO team ads", "budget_id": None,
    }
    assert (rows[4]["brand"], rows[4]["market"]) == ("TBD/All", "Unassigned")
    assert (rows[6]["amount"], rows[6]["currency"], rows[6]["note"]) == (25000.0, "USD", "= 11k cad g ads-")
    assert (rows[7]["amount"], rows[7]["messages"]) == (0.0, ["No approved budget in the sheet, so it's added as 0."])
    assert (rows[8]["amount"], rows[8]["budget_status"], rows[8]["messages"]) == (0.0, "CANCELLED", [])
    assert (rows[9]["amount"], rows[9]["currency"]) == (10000.0, "GBP")
    assert rows[10]["budget_status"] == "NOT_COMMITTED"
    assert rows[11]["status"] == "INVALID"
    assert rows[11]["messages"] == [
        "Unknown brand 'Acme'.", "Unknown market 'Mars'.",
        "Unknown spend progress 'Sleeping'. Use Complete, In Progress, Not Started, Not committed or Canceled.",
    ]
    assert rows[12]["messages"] == ["'Smarch' isn't a month. Use Aug, August or Aug 2026."]
    assert (rows[13]["status"], rows[13]["messages"]) == ("INVALID", ["Same month, brand, market and channel as row 3."])

    # Nothing is saved by a preview
    assert run_sql(f"SELECT COUNT(*) AS n FROM monthly_budgets WHERE {_TEST_MONTHS}")[0]["n"] == 0


def test_confirm_update_and_undo(client, admin, markets):
    # An existing September budget the sheet changes
    existing = client.post("/api/budgets", json={
        "month": f"{YEAR}-09-01", "brand_id": None, "market_id": markets["LONDON"], "amount": 5000, "currency": "USD",
        "status": "APPROVED",
    }, headers=admin["headers"]).json()

    batch_id = _upload(client, admin, _sheet()).json()["batch_id"]
    confirmed = client.post(f"/api/budget-imports/{batch_id}/confirm", headers=admin["headers"])

    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json() | {"batch_id": None} == {
        "batch_id": None, "status": "CONFIRMED", "created": 7, "updated": 1, "unchanged": 0, "notes_added": 7, "skipped": 0,
    }
    september = client.get(f"/api/budgets?month={YEAR}-09-01", headers=admin["headers"]).json()["budgets"]
    by_key = {(b["brand_name"], b["market_name"]): b for b in september}
    assert (by_key[(None, "London")]["amount"], by_key[(None, "London")]["status"]) == (10000, "NOT_COMMITTED")
    assert by_key[("Othership", "London")]["latest_comment"]["body"] == "Campaign paused Sep 21"
    assert (by_key[("Tillerstack", "Germany")]["amount"], by_key[("Tillerstack", "Germany")]["currency"]) == (10000, "GBP")
    assert by_key[("Othership", "North America")]["status"] == "CANCELLED"
    assert len(client.get(f"/api/budgets?month={YEAR}-08-01", headers=admin["headers"]).json()["budgets"]) == 2

    # Uploading the same file again changes nothing
    again = _upload(client, admin, _sheet()).json()
    assert again["counts"] == {"new": 0, "update": 0, "unchanged": 8, "invalid": 3}
    assert [r["budget_id"] for r in again["rows"] if r["row_number"] == 10] == [existing["id"]]

    # A budget changed by hand after the upload is left alone by undo
    workaware = by_key[("WorkAware", "North America")]
    client.patch(f"/api/budgets/{workaware['id']}", json={"amount": 26000}, headers=admin["headers"])

    undone = client.post(f"/api/budget-imports/{batch_id}/undo", headers=admin["headers"])
    assert undone.status_code == 200, undone.text
    assert undone.json() | {"batch_id": None} == {
        "batch_id": None, "status": "UNDONE", "removed": 6, "restored": 1, "notes_removed": 7, "skipped": 1,
    }
    left = client.get(f"/api/budgets?month={YEAR}-09-01", headers=admin["headers"]).json()["budgets"]
    assert sorted((b["brand_name"] or "TBD/All", b["amount"], b["status"]) for b in left) == [
        ("TBD/All", 5000, "APPROVED"), ("WorkAware", 26000, "APPROVED"),
    ]
    assert client.post(f"/api/budget-imports/{batch_id}/undo", headers=admin["headers"]).status_code == 409


def test_only_the_uploader_can_confirm_and_previews_expire(client, admin, other_admin):
    batch_id = _upload(client, admin, _sheet()).json()["batch_id"]

    theirs = client.post(f"/api/budget-imports/{batch_id}/confirm", headers=other_admin["headers"])
    assert (theirs.status_code, theirs.json()["detail"]) == (403, "Only the person who uploaded this file can confirm it.")

    run_sql("UPDATE import_batches SET created_at = now() - interval '25 hours' WHERE id = $1", uuid.UUID(batch_id))
    late = client.post(f"/api/budget-imports/{batch_id}/confirm", headers=admin["headers"])
    assert (late.status_code, late.json()["detail"]) == (410, "This preview has expired. Upload the file again.")
    assert client.post(f"/api/budget-imports/{uuid.uuid4()}/confirm", headers=admin["headers"]).status_code == 404


def test_upload_checks(client, admin, viewer):
    missing = _upload(client, admin, _xlsx([["Month", "Brand", "Notes"], ["Aug", "Othership", "x"]]))
    assert (missing.status_code, missing.json()["detail"]) == (
        422, "The sheet needs Market / Region and Approved Budget columns. Start from the template.",
    )
    assert _upload(client, admin, _sheet(), currency="dollars").status_code == 422
    assert _upload(client, admin, b"not a sheet", "budgets.pdf").status_code == 422
    assert _upload(client, viewer, _sheet()).status_code == 403


def test_template(client, admin):
    response = client.get("/api/budget-imports/template", headers=admin["headers"])

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="budget-tracker-template.xlsx"'
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Budget Tracker"]
    assert [c.value for c in sheet[1]] == TEMPLATE_COLUMNS
    assert "TBD/All" in [c.value for c in workbook["Lists"]["A"]]

    # A filled-in template reads back
    sheet.append(["Oct", "DeskFlex", "Korea", 500, "Not Started"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    preview = _upload(client, admin, buffer.getvalue()).json()
    assert [(r["status"], r["month"], r["amount"]) for r in preview["rows"]] == [("NEW", f"{YEAR}-10-01", 500.0)]
